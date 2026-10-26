import json
import math
import re
import secrets
import bcrypt
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from typing import Optional
from urllib.parse import quote, urlencode, urlsplit, parse_qsl

from fastapi import APIRouter, BackgroundTasks, Depends, Form, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi.templating import Jinja2Templates
from starlette.datastructures import FormData
from sqlalchemy import and_, case, func, or_, text as sql_text
from sqlalchemy.orm import Session, joinedload, selectinload
from sqlalchemy.exc import IntegrityError
from pydantic import BaseModel, Field
from visit_options import OPTION_KINDS, option_names, valid_option
from product_selection import NOT_SELECTED_PRODUCT

from auth import (
    authenticate_user,
    is_super_admin,
    login_user,
    logout_user,
    require_cre_manager_access,
    require_cre_owner_access,
    require_customer_relations_access,
    require_super_admin_access,
    user_store_id,
)
from database import get_db
from dependencies.integration_auth import require_integration_key
from settings import FEEDBACK_AUTO_SEND, TEMPLATE_DIR
from models import CreVisitOption, CreWalkoutFollowup, CustomerProductCategory, AdminUser, CreStaffMember, CustomerFormField, CustomerProductType, CustomerVisit, Store, UserRole, WalkInFeedback
from services.feedback import send_feedback, send_walk_in_feedback_link
from walkout_followups import remember_walkout
from datetime_utils import INDIA_TIME, format_ist
from auth import get_current_user_web


async def _read_form(request: Request):
    # Parse the request body on the event loop before synchronous business work.
    return await request.form()


router = APIRouter()
templates = Jinja2Templates(directory=str(TEMPLATE_DIR))
templates.env.globals["visit_status"] = lambda visit: display_visit_status(visit)
templates.env.filters['ist'] = format_ist

ACTIVITY_TYPES = (
    "New Jewellery",
    "Old Gold",
    "Repair Taken",
    "Repair Delivered",
    "Ordered Delivered",
)
NEW_JEWELLERY_STATUSES = ("Billed", "Walkout")
NEW_JEWELLERY_EDIT_STATUSES = ("Draft", *NEW_JEWELLERY_STATUSES)
VISIT_STATUSES = ("Draft", "Billed", "Walkout", *ACTIVITY_TYPES[1:])
PRODUCT_FREE_STATUSES = ("Old Gold", "Repair Taken", "Repair Delivered")


def display_visit_status(visit):
    if visit.activity_type == "New Jewellery":
        return "Draft" if not visit.status or visit.status == "Draft" else visit.status
    return visit.activity_type


def resolve_visit_status(form):
    value = str(form.get("status", "")).strip()
    # Accept older clients while all current forms use a single Status field.
    if not value:
        activity = str(form.get("activity_type", "")).strip()
        if activity not in ACTIVITY_TYPES:
            raise ValueError("Choose a valid status.")
        value = "Draft" if activity == "New Jewellery" else activity
    if value == "Registered":
        value = "Draft"
    if value not in VISIT_STATUSES:
        raise ValueError("Choose a valid status.")
    if value in ACTIVITY_TYPES[1:]:
        return value, None
    return "New Jewellery", value


def walkout_values(db, form, status, visit=None):
    if status != "Walkout":
        return None, None, None
    remark = str(form.get("remark", "")).strip()
    reason = str(form.get("walkout_reason", "")).strip()
    competitor = str(form.get("competitor_shop", "")).strip()
    if not remark or len(remark) > 1000:
        raise ValueError("Add a walkout remark up to 1000 characters.")
    if not valid_option(db, "reason", reason, visit.walkout_reason if visit else None):
        raise ValueError("Choose a valid walkout reason.")
    if competitor and not valid_option(db, "competitor", competitor, visit.competitor_shop if visit else None):
        raise ValueError("Choose a valid competitor shop.")
    return remark, reason, competitor or None


REPORT_TABS = {"overview", "walkin", "walkout", "competitor", "manager", "sales"}
FIELD_TYPES = ("text", "textarea", "number", "date", "select")
CRE_PANEL_ROLES = {
    UserRole.CUSTOMER_RELATION_EXECUTIVE.value,
    UserRole.SUPER_ADMIN.value,
}
PAGE_SIZE = 20
PAGE_SIZES = (10, 20, 50, 100)
WALKOUT_DEFAULT_FROM = date(2026, 1, 1)


class FeedbackSubmission(BaseModel):
    token: str = Field(min_length=10, max_length=100)
    rating: int = Field(ge=1, le=5)
    comments: Optional[str] = Field(default=None, max_length=2000)


def _effective_store_id(user, requested_store_id=0):
    return requested_store_id if is_super_admin(user) else user_store_id(user)


def _scope_visits(query, user):
    store_id = user_store_id(user)
    return query if store_id is None else query.filter(CustomerVisit.store_id == store_id)


def _scoped_visit(db, visit_id, user):
    return _scope_visits(db.query(CustomerVisit), user).filter(CustomerVisit.id == visit_id).first()


def _staff_query(db, staff_type, store_id=0, active_only=True):
    query = db.query(CreStaffMember).filter(CreStaffMember.staff_type == staff_type)
    if active_only:
        query = query.filter(CreStaffMember.is_active == 1)
    if store_id:
        query = query.filter(CreStaffMember.store_id == store_id)
    return query


def _claim_store_visit_number(db, store_id):
    updated = db.execute(sql_text(
        "UPDATE stores SET next_visit_number=COALESCE(next_visit_number, 1) + 1 "
        "WHERE id=:store_id"
    ), {"store_id": store_id})
    if updated.rowcount != 1:
        raise ValueError("Store was not found.")
    return db.execute(sql_text(
        "SELECT next_visit_number - 1 FROM stores WHERE id=:store_id"
    ), {"store_id": store_id}).scalar_one()


def _pagination_meta(total, page, request, page_key='page', tab=None, per_page=PAGE_SIZE):
    pages = max(1, math.ceil(total / per_page))
    page = min(page, pages)
    offset = (page - 1) * per_page
    first = min(max(1, page - 2), max(1, pages - 4))
    visible = sorted({1, pages, *range(first, min(pages, first + 4) + 1)})
    numbers = []
    for number in visible:
        if numbers and number - numbers[-1] > 1:
            numbers.append(None)
        numbers.append(number)
    params = dict(request.query_params)
    params.pop('message', None)
    params.pop('error', None)
    if tab:
        params['tab'] = tab
    params['per_page'] = per_page
    size_params = {key: value for key, value in params.items()
                   if key not in ('per_page', 'page', 'walkin_page', 'walkout_page', 'competitor_page', 'manager_page', 'sales_page')}
    def url(number):
        return request.url.path + '?' + urlencode({**params, page_key: number})
    return {'page': page, 'pages': pages, 'total': total, 'offset': offset,
            'start': offset + 1 if total else 0, 'end': min(offset + per_page, total),
            'per_page': per_page, 'size_options': PAGE_SIZES,
            'size_params': size_params, 'size_action': request.url.path, 'size_id': page_key,
            'numbers': numbers, 'urls': {number: url(number) for number in visible},
            'previous': url(max(1, page - 1)), 'next': url(min(pages, page + 1))}


def dashboard_day_summary(db, now=None, store_id=0):
    # SQLite CURRENT_TIMESTAMP stores visit timestamps as naive UTC.
    today = (now or datetime.now(INDIA_TIME)).astimezone(INDIA_TIME).date()
    start = datetime.combine(today, time.min, tzinfo=INDIA_TIME)
    end = start + timedelta(days=1)
    daily = db.query(CustomerVisit).filter(
        CustomerVisit.created_at >= start.astimezone(timezone.utc).replace(tzinfo=None),
        CustomerVisit.created_at < end.astimezone(timezone.utc).replace(tzinfo=None),
    )
    if store_id:
        daily = daily.filter(CustomerVisit.store_id == store_id)
    counts = daily.with_entities(
        func.count(CustomerVisit.id),
        func.coalesce(func.sum(case((
            (CustomerVisit.activity_type == 'New Jewellery') &
            or_(CustomerVisit.status == 'Draft', CustomerVisit.status.is_(None)), 1), else_=0)), 0),
        func.coalesce(func.sum(case((
            (CustomerVisit.activity_type == 'New Jewellery') & (CustomerVisit.status == 'Billed'), 1), else_=0)), 0),
        func.coalesce(func.sum(case((
            (CustomerVisit.activity_type == 'New Jewellery') & (CustomerVisit.status == 'Walkout'), 1), else_=0)), 0),
    ).one()
    return {
        'total_visits': counts[0],
        'draft_count': counts[1],
        'billed_count': counts[2],
        'walkout_count': counts[3],
        'summary_date': today.strftime('%d %b %Y'),
        'resets_at': end.isoformat(),
    }


@router.get('/cre/dashboard/summary')
def cre_dashboard_summary(
    store_id: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
):
    from fastapi.responses import JSONResponse
    return JSONResponse(
        dashboard_day_summary(db, store_id=_effective_store_id(current_user, store_id) or 0),
        headers={'Cache-Control': 'no-store'},
    )


@router.get("/admin/customer-relations", include_in_schema=False)
def redirect_legacy_cre_url(request: Request):
    """Move old bookmarks to the dedicated CRE URL while preserving filters."""
    allowed = {
        key: request.query_params[key]
        for key in ("q", "activity", "sales_person", "manager", "page")
        if request.query_params.get(key)
    }
    suffix = f"?{urlencode(allowed)}" if allowed else ""
    return RedirectResponse(f"/cre/dashboard{suffix}", status_code=302)


def _panel_url(message: str = "", error: str = "", path: str = "/cre/dashboard") -> str:
    if error:
        return f"{path}?error={quote(error)}"
    if message:
        return f"{path}?message={quote(message)}"
    return path


def _dashboard_filter_query(q: str, activity: str, sales_person: int, manager: int, store_id: int = 0) -> str:
    values = {}
    if q:
        values["q"] = q
    if activity:
        values["activity"] = activity
    if sales_person:
        values["sales_person"] = sales_person
    if manager:
        values["manager"] = manager
    if store_id:
        values["store_id"] = store_id
    return urlencode(values)


def _report_dates(
    date_from: Optional[str] = Query(default=None),
    date_to: Optional[str] = Query(default=None),
    q: str = Query(default="", max_length=120),
):
    if q.strip():
        return None, None
    dates = []
    for field, value in (('date_from', date_from), ('date_to', date_to)):
        if value is None or not value.strip():
            dates.append(None)
            continue
        try:
            dates.append(date.fromisoformat(value.strip()))
        except ValueError:
            raise HTTPException(status_code=422, detail=[{
                'loc': ['query', field], 'msg': 'Enter a valid date in YYYY-MM-DD format.',
                'type': 'value_error.date',
            }])
    return tuple(dates)


def _report_date_range(date_from: Optional[date], date_to: Optional[date]):
    """Return an inclusive report range, defaulting to the current month."""
    today = datetime.now(INDIA_TIME).date()
    start = date_from or today.replace(day=1)
    end = date_to or today
    if start > end:
        start, end = end, start
    return start, end


def _report_visits_query(
    db: Session,
    start: date,
    end: date,
    store_id: int,
    manager_id: int,
    sales_person_id: int = 0,
    walkout_reason: str = "",
    competitor_shop: str = "",
    q: str = "",
    city: str = "",
):
    query = db.query(CustomerVisit).options(
        joinedload(CustomerVisit.store),
        joinedload(CustomerVisit.handled_by),
        joinedload(CustomerVisit.assigned_manager),
        joinedload(CustomerVisit.product_type),
        joinedload(CustomerVisit.category),
        joinedload(CustomerVisit.walkout_record),
        selectinload(CustomerVisit.followups).joinedload(CreWalkoutFollowup.author),
    )
    if not q.strip():
        query = query.filter(
            CustomerVisit.created_at >= datetime.combine(start, time.min, tzinfo=INDIA_TIME).astimezone(timezone.utc).replace(tzinfo=None),
            CustomerVisit.created_at < datetime.combine(end + timedelta(days=1), time.min, tzinfo=INDIA_TIME).astimezone(timezone.utc).replace(tzinfo=None),
        )
    if store_id:
        query = query.filter(CustomerVisit.store_id == store_id)
    elif city:
        query = query.filter(CustomerVisit.store.has(func.lower(Store.city) == city.lower()))
    if manager_id:
        query = query.filter(CustomerVisit.assigned_manager_id == manager_id)
    if sales_person_id:
        query = query.filter(CustomerVisit.handled_by_id == sales_person_id)
    if walkout_reason:
        query = query.filter(or_(CustomerVisit.walkout_reason == walkout_reason,
                                 CustomerVisit.walkout_record.has(reason=walkout_reason)))
    if competitor_shop:
        query = query.filter(or_(CustomerVisit.competitor_shop == competitor_shop,
                                 CustomerVisit.walkout_record.has(competitor=competitor_shop)))
    if q.strip():
        text = q.strip()
        store_number_match = re.fullmatch(r'([A-Z0-9]{2,8})-(\d{1,19})', text, re.IGNORECASE)
        legacy_number_match = re.fullmatch(r'CRE-(\d{1,19})', text, re.IGNORECASE)
        if legacy_number_match:
            visit_id = int(legacy_number_match.group(1))
            query = query.filter(
                CustomerVisit.id == visit_id if visit_id <= 9223372036854775807 else False
            )
        elif store_number_match:
            sequence = int(store_number_match.group(2))
            if sequence <= 9223372036854775807:
                query = query.filter(and_(
                    CustomerVisit.store_visit_number == sequence,
                    CustomerVisit.store.has(
                        func.lower(Store.visit_prefix) == store_number_match.group(1).lower()
                    ),
                ))
            else:
                query = query.filter(False)
        else:
            query = query.filter(or_(
                CustomerVisit.customer_name.icontains(text, autoescape=True),
                CustomerVisit.mobile_number.icontains(text, autoescape=True),
            ))
    return query


def _build_report_data(
    db: Session,
    start: date,
    end: date,
    store_id: int,
    manager_id: int,
    sales_person_id: int = 0,
    walkout_reason: str = "",
    competitor_shop: str = "",
    q: str = "",
    city: str = "",
):
    visits = _report_visits_query(
        db, start, end, store_id, manager_id, sales_person_id, walkout_reason, competitor_shop, q, city
    ).all()
    billed = [visit for visit in visits if visit.status == "Billed"]
    walkouts = [visit for visit in visits if visit.status == "Walkout"]
    pending = [visit for visit in visits if visit.activity_type == "New Jewellery" and visit.status == "Draft"]
    # The summary must reconcile: every walk-in is either billed, a walkout, or another/pending visitor.
    other_visitors_count = len(visits) - len(billed) - len(walkouts)

    manager_totals = defaultdict(lambda: {"visits": 0, "walkouts": 0, "billed": 0})
    for visit in visits:
        manager_name = visit.assigned_manager.name if visit.assigned_manager else "Not Assigned"
        data = manager_totals[manager_name]
        data["visits"] += 1
        if visit.status == "Walkout":
            data["walkouts"] += 1
        elif visit.status == "Billed":
            data["billed"] += 1
    manager_rows = [
        {
            "name": name,
            **values,
            "conversion_rate": round((values["billed"] / values["visits"] * 100) if values["visits"] else 0, 2),
        }
        for name, values in sorted(manager_totals.items(), key=lambda item: (-item[1]["visits"], item[0].lower()))
    ]

    sales_totals = defaultdict(lambda: {"visits": 0, "walkouts": 0, "billed": 0})
    for visit in visits:
        sales_name = visit.handled_by.name if visit.handled_by else "Not Assigned"
        data = sales_totals[sales_name]
        data["visits"] += 1
        if visit.status == "Walkout":
            data["walkouts"] += 1
        elif visit.status == "Billed":
            data["billed"] += 1
    sales_rows = [
        {
            "name": name,
            **values,
            "conversion_rate": round((values["billed"] / values["visits"] * 100) if values["visits"] else 0, 2),
        }
        for name, values in sorted(sales_totals.items(), key=lambda item: (-item[1]["visits"], item[0].lower()))
    ]

    reason_rows = [
        {"name": name, "count": count}
        for name, count in sorted(
            Counter((visit.walkout_reason or "Not recorded") for visit in walkouts).items(),
            key=lambda item: (-item[1], item[0].lower()),
        )
    ]
    competitor_rows = [
        {"name": name, "count": count}
        for name, count in sorted(
            Counter((visit.competitor_shop or "Not recorded") for visit in walkouts).items(),
            key=lambda item: (-item[1], item[0].lower()),
        )
    ]
    walkin_rows = [
        {
            "date": visit.created_at,
            "visit_number": visit.visit_number,
            "customer_name": visit.customer_name,
            "mobile_number": visit.mobile_number,
            "store_name": visit.store.store_name if visit.store else "Not recorded",
            "product": "" if display_visit_status(visit) in PRODUCT_FREE_STATUSES else (
                visit.product_type.name if visit.product_type else "Not Selected"
            ),
            "activity_type": visit.activity_type,
            "category": "" if display_visit_status(visit) in PRODUCT_FREE_STATUSES else (
                visit.category.name if visit.category else "Not Selected"
            ),
            "status": display_visit_status(visit),
            "sales_person": visit.handled_by.name if visit.handled_by else "Not Assigned",
            "manager": visit.assigned_manager.name if visit.assigned_manager else "Not Assigned",
        }
        for visit in sorted(visits, key=lambda visit: (visit.created_at, visit.id), reverse=True)
    ]
    walkout_rows = [
        {
            "id": visit.id,
            "date": visit.created_at,
            "visit_number": visit.visit_number,
            "customer_name": visit.customer_name,
            "mobile_number": visit.mobile_number,
            "store_name": visit.store.store_name if visit.store else "Not recorded",
            "reason": (visit.walkout_reason if visit.status == 'Walkout' else visit.walkout_record.reason if visit.walkout_record else visit.walkout_reason) or "Not recorded",
            "competitor": (visit.competitor_shop if visit.status == 'Walkout' else visit.walkout_record.competitor if visit.walkout_record else visit.competitor_shop) or "Not recorded",
            "manager": visit.assigned_manager.name if visit.assigned_manager else "Not Assigned",
            "remark": (visit.walkout_record.remark if visit.walkout_record else visit.remark) or "-",
            "status": display_visit_status(visit),
            "product": "" if display_visit_status(visit) in PRODUCT_FREE_STATUSES else (
                visit.product_type.name if visit.product_type else "Not Selected"
            ),
            "category": "" if display_visit_status(visit) in PRODUCT_FREE_STATUSES else (
                visit.category.name if visit.category else "Not Selected"
            ),
            "sales_person": visit.handled_by.name if visit.handled_by else "Not Assigned",
            "followups": {item.slot: item for item in visit.followups},
        }
        for visit in sorted(visits, key=lambda visit: (visit.created_at, visit.id), reverse=True)
        if visit.walkout_record or visit.status == 'Walkout'
    ]
    total = len(visits)
    return {
        "summary": {
            "total": total,
            "billed": len(billed),
            "walkouts": len(walkouts),
            "other_visitors": other_visitors_count,
            "pending": len(pending),
            "billed_percent": round((len(billed) / total * 100) if total else 0, 2),
            "walkout_percent": round((len(walkouts) / total * 100) if total else 0, 2),
            "other_visitors_percent": round((other_visitors_count / total * 100) if total else 0, 2),
        },
        "manager_rows": manager_rows,
        "sales_rows": sales_rows,
        "reason_rows": reason_rows,
        "competitor_rows": competitor_rows,
        "walkin_rows": walkin_rows,
        "walkout_rows": walkout_rows,
        "recovered_count": sum(row['status'] == 'Billed' for row in walkout_rows),
        "register_metrics": {
            "total": len(walkout_rows),
            "open": sum(row['status'] == 'Walkout' for row in walkout_rows),
            "recovered": sum(row['status'] == 'Billed' for row in walkout_rows),
            "pending_followups": sum(row['status'] != 'Billed' and len(row['followups']) < 3 for row in walkout_rows),
        },
    }


def _report_city(db, user, city, store_id):
    if not is_super_admin(user):
        store = db.get(Store, user_store_id(user))
        return store.city if store else ""
    if store_id:
        store = db.get(Store, store_id)
        return store.city if store else ""
    city = " ".join(city.split())
    if not city:
        return ""
    match = db.query(Store.city).filter(func.lower(Store.city) == city.lower()).first()
    return match[0] if match else ""


def _report_staff_query(db, staff_type, store_id, city):
    query = _staff_query(db, staff_type, store_id)
    if city and not store_id:
        query = query.filter(CreStaffMember.store.has(func.lower(Store.city) == city.lower()))
    return query


def _render_staff_page(request, user, staff, stores, selected_store_id, page_title, staff_role, message="", error=""):
    return templates.TemplateResponse("cre_staff.html", {
        "request": request,
        "user": user,
        "user_role": user.role,
        "staff": staff,
        "stores": stores,
        "selected_store_id": selected_store_id,
        "is_super_admin": is_super_admin(user),
        "page_title": page_title,
        "staff_role": staff_role,
        "can_add_staff": True,
        "message": message,
        "error": error,
    })


def _render_cre_page(template_name, request, user, **context):
    return templates.TemplateResponse(template_name, {
        "request": request,
        "user": user,
        "user_role": user.role,
        **context,
    })


def _clean_store_values(store_name, city, store_address, timings):
    return tuple(" ".join(value.split()) for value in (
        store_name, city, store_address, timings
    ))


def _access_url(message="", error=""):
    values = {key: value for key, value in (("message", message), ("error", error)) if value}
    return "/cre/access" + ("?" + urlencode(values) if values else "")


def _product_selection(db, value):
    try:
        product_id = int(value)
    except (TypeError, ValueError):
        product_id = 0
    if product_id:
        product = db.get(CustomerProductType, product_id)
        if product:
            return product
    return db.query(CustomerProductType).filter(
        func.lower(CustomerProductType.name) == NOT_SELECTED_PRODUCT.lower()
    ).first()


@router.get("/cre", include_in_schema=False)
@router.get("/cre/", include_in_schema=False)
def cre_root(request: Request, db: Session = Depends(get_db)):
    user = get_current_user_web(request, db)
    if user and user.role in CRE_PANEL_ROLES:
        return RedirectResponse("/cre/dashboard", status_code=302)
    return RedirectResponse("/cre/login", status_code=302)


@router.get("/cre/login", response_class=HTMLResponse)
def cre_login_page(request: Request, db: Session = Depends(get_db)):
    user = get_current_user_web(request, db)
    if user and user.role in CRE_PANEL_ROLES:
        return RedirectResponse("/cre/dashboard", status_code=302)
    return templates.TemplateResponse("cre_login.html", {"request": request, "error": ""})


@router.post("/cre/login")
def cre_login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    user = authenticate_user(db, username, password)
    if not user or user.role not in CRE_PANEL_ROLES:
        return templates.TemplateResponse(
            "cre_login.html",
            {"request": request, "error": "Invalid CRE username or password."},
            status_code=401,
        )

    request.session.clear()
    login_user(request, user)
    return RedirectResponse("/cre/dashboard", status_code=303)


@router.get("/cre/logout")
def cre_logout(request: Request):
    logout_user(request)
    return RedirectResponse("/cre/login", status_code=303)


@router.get("/cre/dashboard", response_class=HTMLResponse)
def cre_dashboard(
    request: Request,
    q: str = Query(default="", max_length=120),
    activity: str = Query(default="", max_length=40),
    store_id: int = Query(default=0, ge=0),
    sales_person: int = Query(default=0, ge=0),
    manager: int = Query(default=0, ge=0),
    page: int = Query(default=1, ge=1),
    per_page: int = Query(default=PAGE_SIZE, ge=1, le=100),
    message: str = Query(default="", max_length=120),
    error: str = Query(default="", max_length=160),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
):
    per_page = per_page if per_page in PAGE_SIZES else PAGE_SIZE
    selected_store_id = _effective_store_id(current_user, store_id) or 0
    product_types = db.query(CustomerProductType).order_by(
        func.lower(CustomerProductType.name)
    ).all()
    stores_query = db.query(Store)
    if not is_super_admin(current_user):
        stores_query = stores_query.filter(Store.id == selected_store_id)
    stores = stores_query.order_by(func.lower(Store.store_name)).all()
    form_fields = db.query(CustomerFormField).filter(
        CustomerFormField.is_active == 1
    ).order_by(CustomerFormField.id).all()
    form_field_definitions = [
        {"field": field, "options": (field.options or "").splitlines()}
        for field in form_fields
    ]

    sales_people = _staff_query(
        db, "sales_person", selected_store_id
    ).order_by(func.lower(CreStaffMember.name)).all()
    managers = _staff_query(
        db, "manager", selected_store_id
    ).order_by(func.lower(CreStaffMember.name)).all()

    categories = db.query(CustomerProductCategory).order_by(CustomerProductCategory.name).all()
    base_query = db.query(CustomerVisit)
    visits_query = base_query.options(
        joinedload(CustomerVisit.store),
        joinedload(CustomerVisit.product_type),
        joinedload(CustomerVisit.category),
        joinedload(CustomerVisit.handled_by),
        joinedload(CustomerVisit.assigned_manager),
    )
    if selected_store_id:
        visits_query = visits_query.filter(CustomerVisit.store_id == selected_store_id)
    if q.strip():
        search = f"%{q.strip()}%"
        store_number_match = re.fullmatch(r'([A-Z0-9]{2,8})-(\d{1,19})', q.strip(), re.IGNORECASE)
        legacy_number_match = re.fullmatch(r'CRE-(\d{1,19})', q.strip(), re.IGNORECASE)
        if legacy_number_match:
            visit_id = int(legacy_number_match.group(1))
            visits_query = visits_query.filter(
                CustomerVisit.id == visit_id if visit_id <= 9223372036854775807 else False
            )
        elif store_number_match:
            sequence = int(store_number_match.group(2))
            if sequence > 9223372036854775807:
                visits_query = visits_query.filter(False)
            else:
                visits_query = visits_query.filter(
                    CustomerVisit.store_visit_number == sequence,
                    CustomerVisit.store.has(
                        func.lower(Store.visit_prefix) == store_number_match.group(1).lower()
                    ),
                )
        else:
            visits_query = visits_query.join(CustomerVisit.product_type).filter(
                or_(
                    CustomerVisit.customer_name.ilike(search),
                    CustomerVisit.mobile_number.ilike(search),
                    CustomerProductType.name.ilike(search),
                )
            )
    if activity in ("Draft", "Billed", "Walkout"):
        visits_query = visits_query.filter(CustomerVisit.activity_type == "New Jewellery",
                                           CustomerVisit.status == activity)
    elif activity in ACTIVITY_TYPES:
        visits_query = visits_query.filter(CustomerVisit.activity_type == activity)
    if sales_person:
        visits_query = visits_query.filter(CustomerVisit.handled_by_id == sales_person)
    if manager:
        visits_query = visits_query.filter(CustomerVisit.assigned_manager_id == manager)

    total_filtered = visits_query.count()
    total_pages = max(1, math.ceil(total_filtered / per_page))
    page = min(page, total_pages)
    visits = visits_query.order_by(CustomerVisit.created_at.desc(), CustomerVisit.id.desc()).offset(
        (page - 1) * per_page
    ).limit(per_page).all()

    visit_rows = []
    serial_number = (page - 1) * per_page + 1
    for row_index, visit in enumerate(visits):
        try:
            values = json.loads(visit.custom_fields or "{}")
        except (TypeError, ValueError):
            values = {}
        custom_values = [
            {"value": values.get(str(field.id), "")}
            for field in form_fields
        ]
        visit_rows.append({
            "visit": visit,
            "custom_values": custom_values,
            "serial_number": serial_number + row_index,
        })

    daily_summary = dashboard_day_summary(db, store_id=selected_store_id)
    sales_people_count = len(sales_people)

    return _render_cre_page(
        "cre_dashboard.html",
        request,
        current_user,
        product_types=product_types,
        categories=categories,
        stores=stores,
        form_fields=form_fields,
        form_field_definitions=form_field_definitions,
        custom_columns=form_fields,
        sales_people=sales_people,
        managers=managers,
        activity_types=ACTIVITY_TYPES,
        visit_statuses=VISIT_STATUSES,
        new_jewellery_statuses=NEW_JEWELLERY_STATUSES,
        walkout_reasons=option_names(db, "reason"),
        competitor_shops=option_names(db, "competitor"),
        visits=visit_rows,
        query=q,
        selected_activity=activity,
        selected_store_id=selected_store_id,
        selected_sales_person=sales_person,
        selected_manager=manager,
        page=page,
        total_pages=total_pages,
        page_numbers=range(max(1, page - 2), min(total_pages, page + 2) + 1),
        dashboard_pagination=_pagination_meta(total_filtered, page, request, per_page=per_page),
        per_page=per_page,
        total_filtered=total_filtered,
        filter_query=_dashboard_filter_query(
            q, activity, sales_person, manager,
            selected_store_id if is_super_admin(current_user) else 0,
        ),
        **daily_summary,
        sales_people_count=sales_people_count,
        can_manage=is_super_admin(current_user),
        is_super_admin=is_super_admin(current_user),
        can_add_customer=bool(product_types and stores and categories),
        message=message,
        error=error,
    )


@router.get('/cre/walkout-register', response_class=HTMLResponse)
@router.get("/cre/reports", response_class=HTMLResponse)
def cre_reports(
    request: Request,
    dates: tuple = Depends(_report_dates),
    store_id: int = Query(default=0, ge=0),
    city: str = Query(default="", max_length=100),
    manager_id: int = Query(default=0, ge=0),
    sales_person_id: int = Query(default=0, ge=0),
    walkout_reason: str = Query(default="", max_length=80),
    competitor_shop: str = Query(default="", max_length=100),
    q: str = Query(default="", max_length=120),
    tab: str = Query(default="overview", max_length=20),
    walkin_page: int = Query(default=1, ge=1),
    walkout_page: int = Query(default=1, ge=1),
    competitor_page: int = Query(default=1, ge=1),
    manager_page: int = Query(default=1, ge=1),
    sales_page: int = Query(default=1, ge=1),
    per_page: int = Query(default=PAGE_SIZE, ge=1, le=100),
    message: str = Query(default="", max_length=120),
    error: str = Query(default="", max_length=160),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
):
    per_page = per_page if per_page in PAGE_SIZES else PAGE_SIZE
    selected_store_id = _effective_store_id(current_user, store_id) or 0
    selected_city = _report_city(db, current_user, city, selected_store_id)
    date_from, date_to = dates
    standalone = request.url.path == '/cre/walkout-register'
    if standalone and date_from == date(1900, 1, 1):
        date_from = WALKOUT_DEFAULT_FROM
    start, end = _report_date_range(date_from or (WALKOUT_DEFAULT_FROM if standalone else None), date_to)
    active_tab = 'walkout' if standalone else tab if tab in REPORT_TABS else "overview"
    q = q.strip()
    report = _build_report_data(
        db, start, end, selected_store_id, manager_id, sales_person_id,
        walkout_reason, competitor_shop, q, selected_city
    )
    report['competitor_visits'] = [row for row in report['walkout_rows'] if row['competitor'] != 'Not recorded']
    report['manager_page_rows'] = report['manager_rows']
    report['sales_page_rows'] = report['sales_rows']
    report_pages = {}
    for key, selected, rows_key in (
        ('walkin', walkin_page, 'walkin_rows'), ('walkout', walkout_page, 'walkout_rows'),
        ('competitor', competitor_page, 'competitor_visits'),
        ('manager', manager_page, 'manager_page_rows'), ('sales', sales_page, 'sales_page_rows'),
    ):
        rows = report[rows_key]
        paging = _pagination_meta(len(rows), selected, request, key + '_page', key, per_page=per_page)
        report_pages[key] = paging
        report[rows_key] = rows[paging['offset']:paging['offset'] + per_page]
    return _render_cre_page(
        "cre_reports.html",
        request,
        current_user,
        report=report,
        report_pages=report_pages,
        per_page=per_page,
        stores=(db.query(Store) if is_super_admin(current_user) else db.query(Store).filter(
            Store.id == selected_store_id
        )).order_by(func.lower(Store.store_name)).all(),
        cities=sorted({
            value for value, in db.query(Store.city).filter(Store.city.isnot(None)).all() if value
        }, key=str.lower),
        managers=_report_staff_query(
            db, "manager", selected_store_id, selected_city
        ).order_by(func.lower(CreStaffMember.name)).all(),
        sales_people=_report_staff_query(
            db, "sales_person", selected_store_id, selected_city
        ).order_by(func.lower(CreStaffMember.name)).all(),
        date_from=start,
        date_to=end,
        selected_store_id=selected_store_id,
        selected_city=selected_city,
        is_super_admin=is_super_admin(current_user),
        selected_manager_id=manager_id,
        selected_sales_person_id=sales_person_id,
        selected_walkout_reason=walkout_reason,
        selected_competitor_shop=competitor_shop,
        search_query=q,
        clear_search_url=request.url.path + '?' + urlencode({
            key: value for key, value in request.query_params.items()
            if key not in ('q', 'walkout_page', 'walkin_page', 'competitor_page', 'manager_page', 'sales_page')
        }),
        active_tab=active_tab,
        standalone_walkout=standalone,
        today_iso=datetime.now(INDIA_TIME).date().isoformat(),
        message=message,
        error=error,
        report_return_to=('/cre/walkout-register?' if standalone else '/cre/reports?') + urlencode({
            'tab': 'walkout', 'date_from': start.isoformat(), 'date_to': end.isoformat(),
            'city': selected_city, 'store_id': selected_store_id,
            'manager_id': manager_id, 'sales_person_id': sales_person_id,
            'walkout_reason': walkout_reason, 'competitor_shop': competitor_shop,
            'walkout_page': report_pages['walkout']['page'],
            'per_page': per_page, 'q': q,
        }),
        walkout_reasons=option_names(db, "reason", include_history=True),
        competitor_shops=option_names(db, "competitor", include_history=True),
        editable_walkout_reasons=option_names(db, 'reason'),
        editable_competitor_shops=option_names(db, 'competitor'),
    )


def _walkout_excel_response(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Walkout Register'
    sheet.append(['Visit Number', 'Name', 'Mobile', 'Product Type', 'Category Type',
                  'Sales Person', 'Meet by Manager', 'Walkout Reason', 'Status',
                  'Follow up 1', 'Follow up 2', 'Follow up 3'])
    for row in rows:
        values = [row['visit_number'], row['customer_name'], row['mobile_number'],
                  row['product'], row['category'], row['sales_person'], row['manager'],
                  row['reason'], row['status'],
                  *[row['followups'][slot].note if slot in row['followups'] else '' for slot in (1, 2, 3)]]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # All exported values are text, including phone numbers and notes.
            cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.fill = PatternFill('solid', fgColor='E8F5EC' if row['status'] == 'Billed' else 'FFF4F3')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor='AD2635')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for index, width in enumerate((19, 26, 22, 22, 22, 25, 25, 28, 18, 48, 48, 48), 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="cre-walkout-register.xlsx"'})


def _section_excel_response(title, headers, rows, filename):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append([
            "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value
            for value in row
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor='AD2635')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(
            42, max(14, max(len(str(cell.value or '')) for cell in column) + 2)
        )
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'},
    )


def _report_section_excel(report, section, start, end):
    date_suffix = f'{start.isoformat()}-to-{end.isoformat()}'
    if section == 'overview':
        summary = report['summary']
        return _section_excel_response(
            'Overview', ['Metric', 'Count', 'Percentage'], [
                ['Total walk-ins', summary['total'], '100.00%' if summary['total'] else '0.00%'],
                ['Billed', summary['billed'], f"{summary['billed_percent']:.2f}%"],
                ['Other visitors', summary['other_visitors'], f"{summary['other_visitors_percent']:.2f}%"],
                ['Walkouts', summary['walkouts'], f"{summary['walkout_percent']:.2f}%"],
                ['Draft', summary['pending'], ''],
            ], f'cre-overview-{date_suffix}.xlsx'
        )
    if section == 'walkin':
        return _section_excel_response(
            'Walk-in Report',
            ['Date & time (IST)', 'Visit No.', 'Customer', 'Mobile', 'Store', 'Product Type',
             'Category', 'Status', 'Sales Person', 'Manager'],
            [[format_ist(row['date']), row['visit_number'], row['customer_name'], row['mobile_number'],
              row['store_name'], row['product'], row['category'], row['status'], row['sales_person'],
              row['manager']] for row in report['walkin_rows']],
            f'cre-walkin-report-{date_suffix}.xlsx',
        )
    if section == 'walkout':
        return _section_excel_response(
            'Walkout Report',
            ['Date & time (IST)', 'Visit No.', 'Customer', 'Mobile', 'Store', 'Product Type',
             'Category', 'Sales Person', 'Manager', 'Walkout Reason', 'Competitor', 'Remark',
             'Follow-up 1', 'Follow-up 2', 'Follow-up 3', 'Final Status'],
            [[format_ist(row['date']), row['visit_number'], row['customer_name'], row['mobile_number'],
              row['store_name'], row['product'], row['category'], row['sales_person'], row['manager'],
              row['reason'], row['competitor'], row['remark'],
              *[row['followups'][slot].note if slot in row['followups'] else '' for slot in (1, 2, 3)],
              row['status']] for row in report['walkout_rows']],
            f'cre-walkout-report-{date_suffix}.xlsx',
        )
    if section == 'competitor':
        competitor_visits = [row for row in report['walkout_rows'] if row['competitor'] != 'Not recorded']
        return _section_excel_response(
            'Competitor Report',
            ['Date & time (IST)', 'Visit No.', 'Customer', 'Mobile', 'Store', 'Competitor',
             'Walkout Reason', 'Manager'],
            [[format_ist(row['date']), row['visit_number'], row['customer_name'], row['mobile_number'],
              row['store_name'], row['competitor'], row['reason'], row['manager']]
             for row in competitor_visits],
            f'cre-competitor-report-{date_suffix}.xlsx',
        )
    if section in ('manager', 'sales'):
        is_manager = section == 'manager'
        label = 'Manager' if is_manager else 'Sales Person'
        rows = report['manager_rows'] if is_manager else report['sales_rows']
        return _section_excel_response(
            f'{label} Report', [label, 'Total Visits', 'Walkouts', 'Billed', 'Conversion'],
            [[row['name'], row['visits'], row['walkouts'], row['billed'], f"{row['conversion_rate']:.2f}%"]
             for row in rows],
            f"cre-{'manager' if is_manager else 'sales-person'}-report-{date_suffix}.xlsx",
        )
    raise HTTPException(status_code=422, detail='Choose a valid report section.')


@router.get('/cre/walkout-register/export/excel')
@router.get("/cre/reports/export/excel")
def export_cre_report_excel(
    request: Request,
    dates: tuple = Depends(_report_dates),
    store_id: int = Query(default=0, ge=0),
    city: str = Query(default="", max_length=100),
    manager_id: int = Query(default=0, ge=0),
    sales_person_id: int = Query(default=0, ge=0),
    walkout_reason: str = Query(default="", max_length=80),
    competitor_shop: str = Query(default="", max_length=100),
    q: str = Query(default="", max_length=120),
    section: str = Query(default="", max_length=20),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
):
    selected_store_id = _effective_store_id(current_user, store_id) or 0
    selected_city = _report_city(db, current_user, city, selected_store_id)
    date_from, date_to = dates
    standalone = request.url.path == '/cre/walkout-register/export/excel'
    if standalone and date_from == date(1900, 1, 1):
        date_from = WALKOUT_DEFAULT_FROM
    start, end = _report_date_range(date_from or (WALKOUT_DEFAULT_FROM if standalone else None), date_to)
    report = _build_report_data(
        db, start, end, selected_store_id, manager_id, sales_person_id,
        walkout_reason, competitor_shop, q, selected_city
    )
    if standalone:
        return _walkout_excel_response(report['walkout_rows'])
    if section:
        return _report_section_excel(report, section, start, end)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Walk-in Summary"
    summary_sheet.append(["CRE walk-in report", 'Search results across all dates' if q.strip() else f"{start:%d %b %Y} - {end:%d %b %Y}"])
    summary_sheet.append(["Metric", "Count", "Percentage"])
    summary = report["summary"]
    summary_sheet.append(["Billed", summary["billed"], summary["billed_percent"] / 100])
    summary_sheet.append(["Other visitors", summary["other_visitors"], summary["other_visitors_percent"] / 100])
    summary_sheet.append(["Walkouts", summary["walkouts"], summary["walkout_percent"] / 100])
    summary_sheet.append(["Total walk-ins", summary["total"], 1 if summary["total"] else 0])
    for cell in summary_sheet[2]:
        cell.font = cell.font.copy(bold=True)
    for row in (3, 4, 5, 6):
        summary_sheet.cell(row=row, column=3).number_format = "0.00%"
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 16
    summary_sheet.column_dimensions["C"].width = 16

    def add_sheet(title, headers, rows):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for row in rows:
            sheet.append(["'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value for value in row])
        for cell in sheet[1]:
            cell.font = cell.font.copy(bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(32, max(14, max(len(str(cell.value or "")) for cell in column) + 2))

    add_sheet(
        "Manager summary",
        ["Manager", "Visits", "Walkouts", "Billed", "Conversion %"],
        [[row["name"], row["visits"], row["walkouts"], row["billed"], row["conversion_rate"] / 100] for row in report["manager_rows"]],
    )
    manager_sheet = workbook["Manager summary"]
    for cell in manager_sheet["E"][1:]:
        cell.number_format = "0.00%"
    add_sheet(
        "Sales person summary",
        ["Sales person", "Visits", "Walkouts", "Billed", "Conversion %"],
        [[row["name"], row["visits"], row["walkouts"], row["billed"], row["conversion_rate"] / 100] for row in report["sales_rows"]],
    )
    sales_sheet = workbook["Sales person summary"]
    for cell in sales_sheet["E"][1:]:
        cell.number_format = "0.00%"
    add_sheet("Walkout reasons", ["Reason", "Walkouts"], [[row["name"], row["count"]] for row in report["reason_rows"]])
    add_sheet("Competitors", ["Competitor", "Walkouts"], [[row["name"], row["count"]] for row in report["competitor_rows"]])
    add_sheet('Walkout register',
        ['Visit date (IST)', 'Customer', 'Mobile', 'Store', 'Product', 'Sales person', 'Manager',
         'Reason', 'Competitor', 'Remark', 'Follow-up 1', 'Follow-up 2', 'Follow-up 3', 'Final status', 'Category', 'Visit No.'],
        [[format_ist(row['date']), row['customer_name'], row['mobile_number'],
          row['store_name'], row['product'], row['sales_person'], row['manager'], row['reason'],
          row['competitor'], row['remark'],
          *[row['followups'][slot].note if slot in row['followups'] else '' for slot in (1, 2, 3)],
          row['status'], row['category'], row['visit_number']] for row in report['walkout_rows']])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"cre-walkin-report-{start.isoformat()}-to-{end.isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/feedback.html", response_class=HTMLResponse, include_in_schema=False)
def public_feedback_form(request: Request):
    return templates.TemplateResponse("feedback_form.html", {"request": request})


@router.post("/api/feedback", status_code=201)
def submit_feedback(payload: FeedbackSubmission, db: Session = Depends(get_db)):
    feedback = db.query(WalkInFeedback).filter(WalkInFeedback.token == payload.token).first()
    if not feedback:
        raise HTTPException(status_code=404, detail="Feedback link is invalid")
    if feedback.submitted_at is not None:
        raise HTTPException(status_code=409, detail="Feedback has already been submitted")

    feedback.rating = payload.rating
    feedback.comments = (payload.comments or "").strip() or None
    feedback.submitted_at = datetime.utcnow()
    feedback.send_status = "submitted"
    feedback.is_read = False
    feedback.read_at = None
    feedback.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(feedback)
    return {
        "success": True,
        "message": "Thank you. Your feedback has been submitted.",
        "feedback_id": feedback.id,
        "submitted_at": feedback.submitted_at.isoformat(),
    }


@router.get("/cre/feedback", response_class=HTMLResponse)
def feedback_dashboard(
    request: Request,
    status_filter: str = Query(default="all", alias="status", pattern="^(all|submitted|pending)$"),
    q: str = Query(default="", max_length=120),
    page: int = Query(default=1, ge=1),
    per_page: int = Query(default=PAGE_SIZE, ge=1, le=100),
    message: str = Query(default="", max_length=160),
    error: str = Query(default="", max_length=200),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
):
    per_page = per_page if per_page in PAGE_SIZES else PAGE_SIZE
    query = db.query(WalkInFeedback).options(
        joinedload(WalkInFeedback.customer).joinedload(CustomerVisit.store)
    ).join(WalkInFeedback.customer)
    query = _scope_visits(query, current_user)
    if status_filter == "submitted":
        query = query.filter(WalkInFeedback.submitted_at.isnot(None))
    elif status_filter == "pending":
        query = query.filter(WalkInFeedback.submitted_at.is_(None))
    if q.strip():
        term = f"%{q.strip()}%"
        query = query.filter(or_(CustomerVisit.customer_name.ilike(term), CustomerVisit.mobile_number.ilike(term)))

    total = query.count()
    pagination = _pagination_meta(total, page, request, per_page=per_page)
    items = query.order_by(
        WalkInFeedback.is_read.asc(),
        WalkInFeedback.submitted_at.desc(),
        WalkInFeedback.created_at.desc(),
    ).offset(pagination["offset"]).limit(per_page).all()
    return _render_cre_page(
        "cre_feedback.html",
        request,
        current_user,
        feedback_items=items,
        pagination=pagination,
        status_filter=status_filter,
        query=q,
        message=message,
        error=error,
    )


@router.post("/cre/feedback/{feedback_id}/read")
def mark_cre_feedback_read(
    feedback_id: int,
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
):
    feedback = db.query(WalkInFeedback).join(WalkInFeedback.customer).filter(WalkInFeedback.id == feedback_id)
    feedback = _scope_visits(feedback, current_user).first()
    if not feedback:
        raise HTTPException(status_code=404, detail="Feedback not found")
    feedback.is_read = True
    feedback.read_at = datetime.utcnow()
    feedback.updated_at = datetime.utcnow()
    db.commit()
    return RedirectResponse("/cre/feedback", status_code=303)


@router.post("/cre/feedback/{feedback_id}/resend")
def resend_cre_feedback(
    feedback_id: int,
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
):
    query = db.query(WalkInFeedback).options(joinedload(WalkInFeedback.customer)).join(WalkInFeedback.customer)
    feedback = _scope_visits(query, current_user).filter(WalkInFeedback.id == feedback_id).first()
    if not feedback:
        raise HTTPException(status_code=404, detail="Feedback not found")
    try:
        send_feedback(db, feedback, current_user.username, "cre_admin")
    except HTTPException as exc:
        return RedirectResponse(_panel_url(error=exc.detail, path="/cre/feedback"), status_code=303)
    except Exception:
        return RedirectResponse(_panel_url(error="Unable to queue feedback WhatsApp.", path="/cre/feedback"), status_code=303)
    return RedirectResponse(_panel_url(message="Feedback WhatsApp queued.", path="/cre/feedback"), status_code=303)


@router.get(
    "/api/integrations/feedback",
    dependencies=[Depends(require_integration_key)],
)
def list_feedback_for_sync(
    after_time: Optional[datetime] = None,
    after_id: int = Query(default=0, ge=0),
    limit: int = Query(default=100, ge=1, le=500),
    db: Session = Depends(get_db),
):
    if after_time and after_time.tzinfo:
        after_time = after_time.astimezone(timezone.utc).replace(tzinfo=None)
    query = db.query(WalkInFeedback).options(
        joinedload(WalkInFeedback.customer).joinedload(CustomerVisit.store)
    )
    if after_time:
        query = query.filter(or_(
            WalkInFeedback.updated_at > after_time,
            and_(WalkInFeedback.updated_at == after_time, WalkInFeedback.id > after_id),
        ))
    rows = query.order_by(WalkInFeedback.updated_at.asc(), WalkInFeedback.id.asc()).limit(limit).all()
    items = [{
        "id": row.id,
        "walk_in_customer_id": row.customer_visit_id,
        "customer_name": row.customer.customer_name,
        "mobile_number": row.customer.mobile_number,
        "store_id": row.customer.store_id,
        "store_name": row.customer.store.store_name if row.customer.store else None,
        "visit_number": row.customer.visit_number,
        "rating": row.rating,
        "comments": row.comments,
        "send_status": row.send_status,
        "is_read": row.is_read,
        "submitted_at": row.submitted_at.isoformat() if row.submitted_at else None,
        "updated_at": row.updated_at.isoformat(),
        "source_system": "cre",
    } for row in rows]
    return {
        "items": items,
        "count": len(items),
        "next_cursor": {
            "after_time": rows[-1].updated_at.isoformat(),
            "after_id": rows[-1].id,
        } if rows else None,
    }


@router.post(
    "/api/integrations/feedback/{feedback_id}/resend",
    dependencies=[Depends(require_integration_key)],
)
def resend_integration_feedback(feedback_id: int, db: Session = Depends(get_db)):
    feedback = db.query(WalkInFeedback).options(
        joinedload(WalkInFeedback.customer)
    ).filter(WalkInFeedback.id == feedback_id).first()
    if not feedback:
        raise HTTPException(status_code=404, detail="Feedback not found")
    try:
        result = send_feedback(db, feedback, "anand_admin", "anand_admin")
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=502, detail="Unable to queue feedback WhatsApp") from error
    return {
        "success": True,
        "status": "queued",
        "feedback_id": feedback.id,
        "request_id": result["provider_request_id"],
        "sent_to": feedback.sent_to,
        "sent_at": feedback.sent_at.isoformat(),
    }


@router.patch(
    "/api/integrations/feedback/{feedback_id}/read",
    dependencies=[Depends(require_integration_key)],
)
def mark_integration_feedback_read(feedback_id: int, db: Session = Depends(get_db)):
    feedback = db.get(WalkInFeedback, feedback_id)
    if not feedback:
        raise HTTPException(status_code=404, detail="Feedback not found")
    feedback.is_read = True
    feedback.read_at = datetime.utcnow()
    feedback.updated_at = datetime.utcnow()
    db.commit()
    return {"success": True, "feedback_id": feedback.id, "is_read": True}


@router.post("/cre/customers")
def add_customer_visit(
    request: Request,
    background_tasks: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
    form: FormData = Depends(_read_form),
):
    customer_name = " ".join(str(form.get("customer_name", "")).split())
    mobile_number = str(form.get("mobile_number", "")).strip()
    product_type_id = form.get("product_type_id", "")
    try:
        activity_type, status = resolve_visit_status(form)
        remark, walkout_reason, competitor_shop = walkout_values(db, form, status)
    except ValueError as error:
        return RedirectResponse(_panel_url(error=str(error)), status_code=303)
    handled_by_id = form.get("handled_by_id", "")
    assigned_manager_id = form.get("assigned_manager_id", "")
    store_id = form.get("store_id", "")
    digits = re.sub(r"\D", "", mobile_number)

    if not customer_name:
        return RedirectResponse(_panel_url(error="Enter the customer's name."), status_code=303)
    if not re.fullmatch(r"\+?[\d\s().-]+", mobile_number) or len(digits) < 7 or len(digits) > 15:
        return RedirectResponse(_panel_url(error="Enter a valid mobile number."), status_code=303)
    if activity_type not in ACTIVITY_TYPES:
        return RedirectResponse(_panel_url(error="Choose a valid customer activity."), status_code=303)

    product_free = activity_type in PRODUCT_FREE_STATUSES
    product_type = _product_selection(db, "" if product_free else product_type_id)
    if not product_type:
        return RedirectResponse(_panel_url(error="Product selection is unavailable. Refresh and try again."), status_code=303)

    category = None
    if not product_free:
        category_value = str(form.get('category_id', '')).strip()
        if category_value:
            try:
                category = db.get(CustomerProductCategory, int(category_value))
            except (TypeError, ValueError):
                category = None
            if not category:
                return RedirectResponse(_panel_url(error='Choose a valid category.'), status_code=303)

    try:
        handled_by_id = int(handled_by_id) if handled_by_id else None
        assigned_manager_id = int(assigned_manager_id) if assigned_manager_id else None
        store_id = int(store_id) if is_super_admin(current_user) else current_user.store_id
    except (TypeError, ValueError):
        return RedirectResponse(_panel_url(error="Select the store, sales person, and assigned manager."), status_code=303)
    handled_by = db.query(CreStaffMember).filter(
        CreStaffMember.id == handled_by_id,
        CreStaffMember.staff_type == "sales_person",
        CreStaffMember.store_id == store_id,
        CreStaffMember.is_active == 1,
    ).first()
    assigned_manager = db.query(CreStaffMember).filter(
        CreStaffMember.id == assigned_manager_id,
        CreStaffMember.staff_type == "manager",
        CreStaffMember.store_id == store_id,
        CreStaffMember.is_active == 1,
    ).first()
    if (handled_by_id and not handled_by) or (assigned_manager_id and not assigned_manager):
        return RedirectResponse(_panel_url(error="Select an active sales person and manager."), status_code=303)
    store = db.query(Store).filter(Store.id == store_id).first()
    if not store:
        return RedirectResponse(_panel_url(error="Select a valid store."), status_code=303)
    store_visit_number = _claim_store_visit_number(db, store.id)

    form_fields = db.query(CustomerFormField).filter(
        CustomerFormField.is_active == 1
    ).order_by(CustomerFormField.id).all()
    custom_values = {}
    for field in form_fields:
        value = str(form.get(f"custom_{field.id}", "")).strip()
        if field.is_required and not value:
            return RedirectResponse(
                _panel_url(error=f"{field.label} is required."), status_code=303
            )
        if not value:
            continue
        if field.field_type == "select" and value not in (field.options or "").splitlines():
            return RedirectResponse(_panel_url(error=f"Choose a valid value for {field.label}."), status_code=303)
        if field.field_type == "number":
            try:
                float(value)
            except ValueError:
                return RedirectResponse(_panel_url(error=f"Enter a number for {field.label}."), status_code=303)
        if field.field_type == "date":
            try:
                date.fromisoformat(value)
            except ValueError:
                return RedirectResponse(_panel_url(error=f"Enter a valid date for {field.label}."), status_code=303)
        custom_values[str(field.id)] = value

    visit = CustomerVisit(
        customer_name=customer_name[:120],
        mobile_number=mobile_number[:24],
        product_type_id=product_type.id,
        category_id=category.id if category else None,
        handled_by_id=handled_by.id if handled_by else None,
        assigned_manager_id=assigned_manager.id if assigned_manager else None,
        store_id=store.id,
        store_visit_number=store_visit_number,
        activity_type=activity_type,
        status=status,
        remark=remark,
        walkout_reason=walkout_reason,
        competitor_shop=competitor_shop,
        custom_fields=json.dumps(custom_values, ensure_ascii=False),
        created_by=current_user.id,
    )
    db.add(visit)
    remember_walkout(db, visit)
    db.flush()
    feedback = WalkInFeedback(
        customer_visit_id=visit.id,
        token=secrets.token_urlsafe(32),
        send_status="pending",
    )
    db.add(feedback)
    db.commit()
    if FEEDBACK_AUTO_SEND:
        background_tasks.add_task(send_walk_in_feedback_link, feedback.id)
        message = f"Visit {visit.visit_number} saved successfully. Feedback WhatsApp is being queued."
    else:
        message = f"Visit {visit.visit_number} saved successfully. Feedback is ready to send."
    return RedirectResponse(_panel_url(message=message), status_code=303)


@router.post("/cre/customers/{visit_id}/status")
def update_new_jewellery_status(
    visit_id: int,
    status: str = Form(..., max_length=40),
    remark: str = Form(default="", max_length=1000),
    walkout_reason: str = Form(default="", max_length=80),
    competitor_shop: str = Form(default="", max_length=100),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
):
    visit = _scoped_visit(db, visit_id, current_user)
    if not visit:
        return RedirectResponse(_panel_url(error="Customer visit was not found."), status_code=303)
    if visit.status == 'Billed':
        return RedirectResponse(_panel_url(error='Billed entries are locked. Details and status cannot be changed.'), status_code=303)
    try:
        activity_type, stored_status = resolve_visit_status({"status": status})
        remark, walkout_reason, competitor_shop = walkout_values(db, {
            "remark": remark, "walkout_reason": walkout_reason,
            "competitor_shop": competitor_shop}, stored_status, visit)
    except ValueError as error:
        return RedirectResponse(_panel_url(error=str(error)), status_code=303)
    remember_walkout(db, visit)
    visit.activity_type = activity_type
    visit.status = stored_status
    if activity_type in PRODUCT_FREE_STATUSES:
        product_type = _product_selection(db, "")
        if not product_type:
            return RedirectResponse(_panel_url(error="Product selection is unavailable. Refresh and try again."), status_code=303)
        visit.product_type_id = product_type.id
        visit.category_id = None
    visit.remark = remark
    visit.walkout_reason = walkout_reason
    visit.competitor_shop = competitor_shop
    remember_walkout(db, visit)
    if not _commit_visit_change(db):
        return RedirectResponse(_panel_url(error='Billed entries are locked. Refresh the page.'), status_code=303)
    return RedirectResponse(_panel_url(message=f"Status updated to {status}."), status_code=303)


@router.post("/cre/customers/{visit_id}/edit")
def edit_customer_visit(
    visit_id: int,
    request: Request,
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
    form: FormData = Depends(_read_form),
):
    visit = _scoped_visit(db, visit_id, current_user)
    if not visit:
        return RedirectResponse(_panel_url(error="Customer visit was not found."), status_code=303)

    if visit.status == 'Billed':
        return RedirectResponse(_panel_url(error='Billed entries are locked. Details and status cannot be changed.'), status_code=303)
    customer_name = " ".join(str(form.get("customer_name", "")).split())
    mobile_number = str(form.get("mobile_number", "")).strip()
    digits = re.sub(r"\D", "", mobile_number)
    if not customer_name or len(customer_name) > 120:
        return RedirectResponse(_panel_url(error="Enter a customer name up to 120 characters."), status_code=303)
    if len(mobile_number) > 24 or not re.fullmatch(r"\+?[\d\s().-]+", mobile_number) or not 7 <= len(digits) <= 15:
        return RedirectResponse(_panel_url(error="Enter a valid mobile number."), status_code=303)

    try:
        activity_type, status = resolve_visit_status(form)
    except ValueError as error:
        return RedirectResponse(_panel_url(error=str(error)), status_code=303)

    try:
        handled_by_id = int(form.get("handled_by_id") or 0)
        assigned_manager_id = int(form.get("assigned_manager_id") or 0)
    except (TypeError, ValueError):
        return RedirectResponse(_panel_url(error="Choose a store, sales person, and manager."), status_code=303)

    submitted_store_id = form.get("store_id")
    if submitted_store_id not in (None, ""):
        try:
            if int(submitted_store_id) != visit.store_id:
                return RedirectResponse(_panel_url(error="A visit cannot be moved to another store."), status_code=303)
        except (TypeError, ValueError):
            return RedirectResponse(_panel_url(error="Choose a valid store."), status_code=303)
    store_id = visit.store_id

    product_free = activity_type in PRODUCT_FREE_STATUSES
    product_type = _product_selection(db, "" if product_free else form.get("product_type_id", ""))
    handled_by = db.query(CreStaffMember).filter(
        CreStaffMember.id == handled_by_id,
        CreStaffMember.staff_type == "sales_person",
        CreStaffMember.store_id == store_id,
        or_(CreStaffMember.is_active == 1, CreStaffMember.id == visit.handled_by_id),
    ).first()
    assigned_manager = db.query(CreStaffMember).filter(
        CreStaffMember.id == assigned_manager_id,
        CreStaffMember.staff_type == "manager",
        CreStaffMember.store_id == store_id,
        or_(CreStaffMember.is_active == 1, CreStaffMember.id == visit.assigned_manager_id),
    ).first()
    store = db.query(Store).filter(Store.id == store_id).first()
    if not product_type or not store or (handled_by_id and not handled_by) or (assigned_manager_id and not assigned_manager):
        return RedirectResponse(_panel_url(error="Choose a valid store, sales person, and manager."), status_code=303)

    remark = str(form.get("remark", "")).strip()
    if len(remark) > 1000:
        return RedirectResponse(_panel_url(error="Remark must be 1000 characters or fewer."), status_code=303)
    category = None
    if not product_free:
        category_value = str(form.get('category_id', '')).strip()
        if category_value:
            try:
                category = db.get(CustomerProductCategory, int(category_value))
            except (TypeError, ValueError):
                category = None
            if not category:
                return RedirectResponse(_panel_url(error='Choose a valid category.'), status_code=303)

    try:
        walkout_remark, walkout_reason, competitor_shop = walkout_values(db, form, status, visit)
        if status == "Walkout":
            remark = walkout_remark
    except ValueError as error:
        return RedirectResponse(_panel_url(error=str(error)), status_code=303)

    try:
        custom_values = json.loads(visit.custom_fields or "{}")
        if not isinstance(custom_values, dict):
            custom_values = {}
    except (TypeError, ValueError):
        custom_values = {}

    form_fields = db.query(CustomerFormField).filter(
        CustomerFormField.is_active == 1
    ).order_by(CustomerFormField.id).all()
    for field in form_fields:
        value = str(form.get(f"custom_{field.id}", "")).strip()
        if field.is_required and not value:
            return RedirectResponse(_panel_url(error=f"{field.label} is required."), status_code=303)
        if len(value) > 500:
            return RedirectResponse(_panel_url(error=f"{field.label} must be 500 characters or fewer."), status_code=303)
        if field.field_type == "select" and value and value not in (field.options or "").splitlines():
            return RedirectResponse(_panel_url(error=f"Choose a valid value for {field.label}."), status_code=303)
        if field.field_type == "number" and value:
            try:
                if not math.isfinite(float(value)):
                    raise ValueError
            except ValueError:
                return RedirectResponse(_panel_url(error=f"Enter a valid number for {field.label}."), status_code=303)
        if field.field_type == "date" and value:
            try:
                date.fromisoformat(value)
            except ValueError:
                return RedirectResponse(_panel_url(error=f"Enter a valid date for {field.label}."), status_code=303)
        if value:
            custom_values[str(field.id)] = value
        else:
            custom_values.pop(str(field.id), None)

    remember_walkout(db, visit)
    visit.customer_name = customer_name
    visit.mobile_number = mobile_number
    visit.product_type_id = product_type.id
    visit.category_id = category.id if category else None
    visit.handled_by_id = handled_by.id if handled_by else None
    visit.assigned_manager_id = assigned_manager.id if assigned_manager else None
    visit.store_id = store.id
    visit.activity_type = activity_type
    visit.status = status
    visit.remark = remark or None
    visit.walkout_reason = walkout_reason or None
    visit.competitor_shop = competitor_shop or None
    visit.custom_fields = json.dumps(custom_values, ensure_ascii=False)
    remember_walkout(db, visit)
    if not _commit_visit_change(db):
        return RedirectResponse(_panel_url(error='Billed entries are locked. Refresh the page.'), status_code=303)
    return RedirectResponse(_panel_url(message="Customer visit updated."), status_code=303)


def _walkout_report_redirect(form, message='', error=''):
    target = urlsplit(str(form.get('return_to', '')))
    allowed = {'date_from', 'date_to', 'store_id', 'manager_id', 'sales_person_id',
               'walkout_reason', 'competitor_shop', 'walkout_page', 'per_page', 'q'}
    params = {key: value for key, value in parse_qsl(target.query)
              if key in allowed} if target.path in ('/cre/reports', '/cre/walkout-register') and not target.netloc else {}
    destination = '/cre/walkout-register' if target.path == '/cre/walkout-register' and not target.netloc else '/cre/reports'
    if destination == '/cre/reports':
        params.update(tab='walkout')
    if message:
        params['message'] = message
    if error:
        params['error'] = error
    return RedirectResponse(destination + '?' + urlencode(params), status_code=303)


def _commit_visit_change(db):
    try:
        db.commit()
        return True
    except IntegrityError:
        db.rollback()
        return False


@router.post('/cre/reports/walkouts/{visit_id}/followups/{slot}')
def save_walkout_followup(
    visit_id: int, slot: int, request: Request,
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
    form: FormData = Depends(_read_form),
):
    visit = _scoped_visit(db, visit_id, current_user)
    if not visit or not (visit.walkout_record or visit.status == 'Walkout'):
        return _walkout_report_redirect(form, error='Walkout visit was not found.')
    if visit.status == 'Billed':
        return _walkout_report_redirect(form, error='Billed entries are locked. No further changes are allowed.')
    note = str(form.get('note', '')).strip()
    if slot not in (1, 2, 3) or not note or len(note) > 1000:
        return _walkout_report_redirect(form, error='Add a follow-up note up to 1000 characters.')
    saved = {item.slot for item in visit.followups}
    if slot in saved:
        return _walkout_report_redirect(form, error='This follow-up is already saved and locked.')
    if any(previous not in saved for previous in range(1, slot)):
        return _walkout_report_redirect(form, error='Save the earlier follow-up first.')
    remember_walkout(db, visit)
    db.add(CreWalkoutFollowup(visit_id=visit.id, slot=slot, note=note, created_by=current_user.id))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return _walkout_report_redirect(form, error='This follow-up is already saved and locked.')
    return _walkout_report_redirect(form, message=f'Follow-up {slot} saved and locked.')


@router.post('/cre/reports/walkouts/{visit_id}/edit')
@router.post('/cre/reports/walkouts/{visit_id}/outcome')
def update_walkout_outcome(
    visit_id: int, request: Request, db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_customer_relations_access),
    form: FormData = Depends(_read_form),
):
    visit = _scoped_visit(db, visit_id, current_user)
    if not visit or not (visit.walkout_record or visit.status == 'Walkout'):
        return _walkout_report_redirect(form, error='Walkout visit was not found.')
    if visit.status == 'Billed':
        return _walkout_report_redirect(form, error='Billed entries are locked. No further changes are allowed.')
    if set(form) - {'status', 'walkout_reason', 'competitor_shop', 'return_to'}:
        return _walkout_report_redirect(form, error='Only status, walkout reason and competitor can be edited here.')
    status = str(form.get('status', ''))
    if status not in ('Walkout', 'Billed'):
        return _walkout_report_redirect(form, error='Choose Walkout or Billed as the final status.')
    remember_walkout(db, visit)
    previous_reason = visit.walkout_reason or (visit.walkout_record.reason if visit.walkout_record else '')
    previous_competitor = visit.competitor_shop or (visit.walkout_record.competitor if visit.walkout_record else '')
    editing = request.url.path.endswith('/edit')
    reason = str(form.get('walkout_reason', '' if editing else previous_reason) or '').strip()
    competitor = str(form.get('competitor_shop', '' if editing else previous_competitor) or '').strip()
    if (editing or 'walkout_reason' in form) and not valid_option(db, 'reason', reason, previous_reason):
        return _walkout_report_redirect(form, error='Choose a valid walkout reason.')
    if competitor and not valid_option(db, 'competitor', competitor, previous_competitor):
        return _walkout_report_redirect(form, error='Choose a valid competitor shop.')
    visit.walkout_reason = reason or None
    visit.competitor_shop = competitor or None
    if visit.walkout_record:
        visit.walkout_record.reason = reason or None
        visit.walkout_record.competitor = competitor or None
    visit.activity_type = 'New Jewellery'
    visit.status = status
    if not _commit_visit_change(db):
        return _walkout_report_redirect(form, error='Billed entries are locked. Refresh the page.')
    return _walkout_report_redirect(form, message=f'Walkout details saved. Status: {status}.')


def _create_staff(db: Session, name: str, staff_type: str, store_id: int, created_by: int):
    name = " ".join(name.split())
    if not name:
        return "Enter the staff member's name."
    if not db.get(Store, store_id):
        return "Choose a valid store."
    if db.query(CreStaffMember).filter(
        CreStaffMember.staff_type == staff_type,
        CreStaffMember.store_id == store_id,
        func.lower(CreStaffMember.name) == name.lower(),
    ).first():
        return "That name is already in the list."

    db.add(CreStaffMember(
        name=name[:100], staff_type=staff_type, store_id=store_id, created_by=created_by
    ))
    db.commit()
    return ""


@router.get("/cre/sales-people", response_class=HTMLResponse)
def sales_people_page(
    request: Request,
    store_id: int = Query(default=0, ge=0),
    message: str = Query(default="", max_length=120),
    error: str = Query(default="", max_length=160),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    selected_store_id = _effective_store_id(current_user, store_id) or 0
    staff = _staff_query(db, "sales_person", selected_store_id).order_by(
        func.lower(CreStaffMember.name)
    ).all()
    stores = (db.query(Store) if is_super_admin(current_user) else db.query(Store).filter(
        Store.id == selected_store_id
    )).order_by(func.lower(Store.store_name)).all()
    return _render_staff_page(request, current_user, staff, stores, selected_store_id, "Sales People", "sales_person", message, error)


@router.post("/cre/sales-people")
def add_sales_person(
    name: str = Form(..., max_length=100),
    store_id: int = Form(default=0, ge=0),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    selected_store_id = _effective_store_id(current_user, store_id) or 0
    error = _create_staff(db, name, "sales_person", selected_store_id, current_user.id)
    if error:
        return RedirectResponse(_panel_url(error=error, path="/cre/sales-people"), status_code=303)
    return RedirectResponse(_panel_url(message="Sales person added to the list.", path="/cre/sales-people"), status_code=303)


@router.get("/cre/managers", response_class=HTMLResponse)
def managers_page(
    request: Request,
    store_id: int = Query(default=0, ge=0),
    message: str = Query(default="", max_length=120),
    error: str = Query(default="", max_length=160),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_owner_access),
):
    selected_store_id = _effective_store_id(current_user, store_id) or 0
    staff = _staff_query(db, "manager", selected_store_id).order_by(
        func.lower(CreStaffMember.name)
    ).all()
    stores = (db.query(Store) if is_super_admin(current_user) else db.query(Store).filter(
        Store.id == selected_store_id
    )).order_by(func.lower(Store.store_name)).all()
    return _render_staff_page(request, current_user, staff, stores, selected_store_id, "Managers", "manager", message, error)


@router.post("/cre/managers")
def add_manager(
    name: str = Form(..., max_length=100),
    store_id: int = Form(default=0, ge=0),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_owner_access),
):
    selected_store_id = _effective_store_id(current_user, store_id) or 0
    error = _create_staff(db, name, "manager", selected_store_id, current_user.id)
    if error:
        return RedirectResponse(_panel_url(error=error, path="/cre/managers"), status_code=303)
    return RedirectResponse(_panel_url(message="Manager added to the list.", path="/cre/managers"), status_code=303)


@router.get("/cre/access", response_class=HTMLResponse)
def access_management_page(
    request: Request,
    message: str = Query(default="", max_length=160),
    error: str = Query(default="", max_length=200),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_super_admin_access),
):
    return _render_cre_page(
        "cre_access.html", request, current_user,
        stores=db.query(Store).order_by(func.lower(Store.city), func.lower(Store.store_name)).all(),
        accounts=db.query(AdminUser).filter(
            AdminUser.role == UserRole.CUSTOMER_RELATION_EXECUTIVE.value
        ).order_by(func.lower(AdminUser.username)).all(),
        message=message,
        error=error,
    )


@router.post("/cre/access/stores")
def create_store(
    store_name: str = Form(..., max_length=160),
    city: str = Form(..., max_length=100),
    visit_prefix: str = Form(..., max_length=8),
    store_address: str = Form(..., max_length=500),
    timings: str = Form(..., max_length=160),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_super_admin_access),
):
    store_name, city, store_address, timings = _clean_store_values(
        store_name, city, store_address, timings
    )
    visit_prefix = visit_prefix.strip().upper()
    if not all((store_name, city, store_address, timings)):
        return RedirectResponse(_access_url(error="Complete all store fields."), status_code=303)
    if not re.fullmatch(r"[A-Z0-9]{2,8}", visit_prefix):
        return RedirectResponse(_access_url(error="Visit prefix must be 2-8 letters or numbers."), status_code=303)
    if db.query(Store).filter(func.lower(Store.store_name) == store_name.lower()).first():
        return RedirectResponse(_access_url(error="A store with this name already exists."), status_code=303)
    if db.query(Store).filter(func.lower(Store.visit_prefix) == visit_prefix.lower()).first():
        return RedirectResponse(_access_url(error="That visit prefix is already assigned."), status_code=303)
    db.add(Store(
        store_name=store_name, city=city, visit_prefix=visit_prefix,
        next_visit_number=1, store_address=store_address, timings=timings
    ))
    db.commit()
    return RedirectResponse(_access_url(message="Store created successfully."), status_code=303)


@router.post("/cre/access/stores/{store_id}/update")
def update_store(
    store_id: int,
    store_name: str = Form(..., max_length=160),
    city: str = Form(..., max_length=100),
    visit_prefix: str = Form(..., max_length=8),
    store_address: str = Form(..., max_length=500),
    timings: str = Form(..., max_length=160),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_super_admin_access),
):
    store = db.get(Store, store_id)
    if not store:
        return RedirectResponse(_access_url(error="Store was not found."), status_code=303)
    store_name, city, store_address, timings = _clean_store_values(
        store_name, city, store_address, timings
    )
    visit_prefix = visit_prefix.strip().upper()
    if not all((store_name, city, store_address, timings)):
        return RedirectResponse(_access_url(error="Complete all store fields."), status_code=303)
    if visit_prefix != store.visit_prefix:
        return RedirectResponse(_access_url(error="Visit prefix cannot be changed after store creation."), status_code=303)
    duplicate = db.query(Store).filter(
        Store.id != store_id, func.lower(Store.store_name) == store_name.lower()
    ).first()
    if duplicate:
        return RedirectResponse(_access_url(error="A store with this name already exists."), status_code=303)
    store.store_name = store_name
    store.city = city
    store.store_address = store_address
    store.timings = timings
    db.commit()
    return RedirectResponse(_access_url(message="Store details updated."), status_code=303)


@router.post("/cre/access/accounts")
def create_store_account(
    username: str = Form(..., max_length=50),
    display_name: str = Form(default="", max_length=100),
    password: str = Form(..., max_length=128),
    store_id: int = Form(..., ge=1),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_super_admin_access),
):
    username = username.strip()
    display_name = " ".join(display_name.split())
    if not re.fullmatch(r"[A-Za-z0-9_.-]{3,50}", username):
        return RedirectResponse(_access_url(error="Username must be 3-50 letters, numbers, dots, dashes or underscores."), status_code=303)
    if len(password) < 8:
        return RedirectResponse(_access_url(error="Password must contain at least 8 characters."), status_code=303)
    if not db.get(Store, store_id):
        return RedirectResponse(_access_url(error="Choose a valid store."), status_code=303)
    if db.query(AdminUser).filter(func.lower(AdminUser.username) == username.lower()).first():
        return RedirectResponse(_access_url(error="That username already exists."), status_code=303)
    db.add(AdminUser(
        username=username,
        display_name=display_name or None,
        password_hash=bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8"),
        role=UserRole.CUSTOMER_RELATION_EXECUTIVE.value,
        store_id=store_id,
    ))
    db.commit()
    return RedirectResponse(_access_url(message="Store login created successfully."), status_code=303)


@router.post("/cre/access/accounts/{account_id}/update")
def update_store_account(
    account_id: int,
    display_name: str = Form(default="", max_length=100),
    password: str = Form(default="", max_length=128),
    store_id: int = Form(..., ge=1),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_super_admin_access),
):
    account = db.query(AdminUser).filter(
        AdminUser.id == account_id,
        AdminUser.role == UserRole.CUSTOMER_RELATION_EXECUTIVE.value,
    ).first()
    if not account:
        return RedirectResponse(_access_url(error="Store login was not found."), status_code=303)
    if not db.get(Store, store_id):
        return RedirectResponse(_access_url(error="Choose a valid store."), status_code=303)
    if password and len(password) < 8:
        return RedirectResponse(_access_url(error="New password must contain at least 8 characters."), status_code=303)
    account.display_name = " ".join(display_name.split()) or None
    account.store_id = store_id
    if password:
        account.password_hash = bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")
    db.commit()
    return RedirectResponse(_access_url(message="Store login updated."), status_code=303)


@router.get("/cre/products", response_class=HTMLResponse)
def products_page(
    request: Request,
    message: str = Query(default="", max_length=120),
    error: str = Query(default="", max_length=160),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    products = db.query(CustomerProductType).order_by(func.lower(CustomerProductType.name)).all()
    return _render_cre_page(
        "cre_products.html", request, current_user,
        products=products, categories=db.query(CustomerProductCategory).order_by(CustomerProductCategory.name).all(), message=message, error=error,
    )


@router.post("/cre/products")
def add_product(
    name: str = Form(..., max_length=80),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    name = " ".join(name.split())
    path = "/cre/products"
    if not name:
        return RedirectResponse(_panel_url(error="Enter a product name.", path=path), status_code=303)
    if db.query(CustomerProductType).filter(func.lower(CustomerProductType.name) == name.lower()).first():
        return RedirectResponse(_panel_url(error="That product already exists.", path=path), status_code=303)
    db.add(CustomerProductType(name=name, created_by=current_user.id))
    db.commit()
    return RedirectResponse(_panel_url(message="Product added to the customer form.", path=path), status_code=303)


@router.post('/cre/categories')
def add_product_category(
    name: str = Form(..., max_length=80), db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    name = ' '.join(name.split())
    if not name:
        return RedirectResponse(_panel_url(error='Enter a category name.', path='/cre/products'), status_code=303)
    db.add(CustomerProductCategory(name=name, created_by=current_user.id))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse(_panel_url(error='That category already exists.', path='/cre/products'), status_code=303)
    return RedirectResponse(_panel_url(message='Category added to the customer form.', path='/cre/products'), status_code=303)


@router.get("/cre/fields", response_class=HTMLResponse)
def fields_page(
    request: Request,
    message: str = Query(default="", max_length=120),
    error: str = Query(default="", max_length=160),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    fields = db.query(CustomerFormField).order_by(CustomerFormField.id).all()
    return _render_cre_page(
        "cre_fields.html", request, current_user,
        fields=fields,
        active_field_count=sum(1 for field in fields if field.is_active),
        field_types=FIELD_TYPES,
        message=message,
        error=error,
    )


@router.post("/cre/fields")
def add_form_field(
    label: str = Form(..., max_length=80),
    field_type: str = Form(..., max_length=20),
    options: str = Form(default="", max_length=1000),
    is_required: str = Form(default=""),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    label = " ".join(label.split())
    options = "\n".join(dict.fromkeys(
        " ".join(value.split()) for value in options.split(",") if value.strip()
    ))
    path = "/cre/fields"
    if not label:
        return RedirectResponse(_panel_url(error="Enter a field name.", path=path), status_code=303)
    if field_type not in FIELD_TYPES:
        return RedirectResponse(_panel_url(error="Choose a valid field type.", path=path), status_code=303)
    if field_type == "select" and not options:
        return RedirectResponse(_panel_url(error="Add comma-separated options for this dropdown field.", path=path), status_code=303)
    if db.query(CustomerFormField).filter(func.lower(CustomerFormField.label) == label.lower()).first():
        return RedirectResponse(_panel_url(error="A field with that name already exists.", path=path), status_code=303)

    db.add(CustomerFormField(
        label=label,
        field_type=field_type,
        options=options or None,
        is_required=1 if is_required == "on" else 0,
        created_by=current_user.id,
    ))
    db.commit()
    return RedirectResponse(_panel_url(message="Field added to the customer form.", path=path), status_code=303)


@router.post("/cre/fields/{field_id}/remove")
def remove_form_field(
    field_id: int,
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    path = "/cre/fields"
    field = db.query(CustomerFormField).filter(CustomerFormField.id == field_id).first()
    if not field:
        return RedirectResponse(_panel_url(error="Form field was not found.", path=path), status_code=303)
    if not field.is_active:
        return RedirectResponse(_panel_url(message="This field is already removed.", path=path), status_code=303)

    field.is_active = 0
    db.commit()
    return RedirectResponse(
        _panel_url(message="Field removed from new forms. Existing customer data is preserved.", path=path),
        status_code=303,
    )


@router.post("/cre/fields/{field_id}/restore")
def restore_form_field(
    field_id: int,
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    path = "/cre/fields"
    field = db.query(CustomerFormField).filter(CustomerFormField.id == field_id).first()
    if not field:
        return RedirectResponse(_panel_url(error="Form field was not found.", path=path), status_code=303)
    if field.is_active:
        return RedirectResponse(_panel_url(message="This field is already active.", path=path), status_code=303)

    field.is_active = 1
    db.commit()
    return RedirectResponse(
        _panel_url(message="Field restored to the customer form.", path=path),
        status_code=303,
    )


@router.get("/cre/walkout-options", response_class=HTMLResponse)
def walkout_options_page(
    request: Request, message: str = "", error: str = "",
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    groups = [dict(kind=kind, title=values[0], max_length=values[1],
                   options=db.query(CreVisitOption).filter_by(kind=kind)
                   .order_by(CreVisitOption.is_active.desc(), CreVisitOption.name).all())
              for kind, values in OPTION_KINDS.items()]
    return templates.TemplateResponse("cre_walkout_options.html", dict(
        request=request, user=current_user, user_role=current_user.role,
        groups=groups, message=message, error=error))


@router.post("/cre/walkout-options/{kind}")
def add_walkout_option(
    kind: str, name: str = Form(..., max_length=100),
    db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    path = "/cre/walkout-options"
    name = " ".join(name.split())
    if kind not in OPTION_KINDS or not name or len(name) > OPTION_KINDS[kind][1]:
        return RedirectResponse(_panel_url(error="Enter a valid name within the character limit.", path=path), status_code=303)
    existing = db.query(CreVisitOption).filter_by(kind=kind, name=name).first()
    if existing:
        return RedirectResponse(_panel_url(error="This name already exists. Restore it if removed.", path=path), status_code=303)
    db.add(CreVisitOption(kind=kind, name=name, created_by=current_user.id))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse(_panel_url(error="This name already exists.", path=path), status_code=303)
    return RedirectResponse(_panel_url(message="Option added.", path=path), status_code=303)


@router.post("/cre/walkout-options/{option_id}/{action}")
def toggle_walkout_option(
    option_id: int, action: str, db: Session = Depends(get_db),
    current_user: AdminUser = Depends(require_cre_manager_access),
):
    path = "/cre/walkout-options"
    option = db.get(CreVisitOption, option_id)
    if not option or action not in ("remove", "restore"):
        return RedirectResponse(_panel_url(error="Option or action was not found.", path=path), status_code=303)
    option.is_active = int(action == "restore")
    db.commit()
    return RedirectResponse(_panel_url(message="Option restored." if option.is_active else "Option removed from new selections. Existing visits are unchanged.", path=path), status_code=303)
