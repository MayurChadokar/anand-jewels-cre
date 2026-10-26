import hashlib
from contextlib import closing
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
import io
import json
import os
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
from datetime import datetime, timezone
from http.cookiejar import CookieJar
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import quote, urlencode, urlparse
from urllib.request import build_opener, HTTPCookieProcessor, Request

import bcrypt
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from database import Base
from models import AdminUser, Store, CustomerProductType, CustomerFormField, CreStaffMember, CustomerVisit, CustomerProductCategory, WalkInFeedback
from migrate_from_main import migrate
from migrate_store_database import migrate as migrate_store_database
ROOT=Path(__file__).resolve().parents[1]
PASSWORD='Standalone-test-password'

def fixture(path):
    engine=create_engine('sqlite:///'+Path(path).as_posix())
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        db.add(Store(id=1,store_name='Test Store',city='Indore',store_address='Test address',timings='10-6'))
        db.commit()
        hashed=bcrypt.hashpw(PASSWORD.encode(),bcrypt.gensalt()).decode()
        db.add_all([AdminUser(id=1,username='owner',role='super_admin',password_hash=hashed),
                    AdminUser(id=2,username='historical',role='contact_manager',password_hash=hashed),
                    AdminUser(id=3,username='unrelated',role='store_admin',password_hash=hashed),
                    AdminUser(id=4,username='executive',role='customer_relation_executive',store_id=1,password_hash=hashed)])
        db.commit()
        db.add_all([CustomerProductType(id=1,name='Ring',created_by=1),
                    CustomerProductCategory(id=1,name='Gold',created_by=1),
                    CustomerFormField(id=1,label='Budget',field_type='number',is_required=1,created_by=1),
                    CreStaffMember(id=1,name='Sales Test',staff_type='sales_person',store_id=1,created_by=2),
                    CreStaffMember(id=2,name='Manager Test',staff_type='manager',store_id=1,created_by=1)])
        db.commit()
        db.add(CustomerVisit(id=1,customer_name='Initial',mobile_number='9876543210',product_type_id=1,
                             handled_by_id=1,assigned_manager_id=2,store_id=1,activity_type='New Jewellery',
                             status='Draft',created_by=4,custom_fields='{"1": "5000"}'))
        db.commit()
    engine.dispose()

class MigrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.source=Path(self.tmp.name)/'source.db'
        self.target=Path(self.tmp.name)/'cre.db'
        fixture(self.source)

    def test_consistent_wal_snapshot_and_only_cre_data(self):
        with closing(sqlite3.connect(str(self.source))) as live:
            live.execute('PRAGMA journal_mode=WAL')
            live.execute("UPDATE customer_visits SET customer_name='Committed WAL' WHERE id=1")
            live.execute('CREATE TABLE unrelated_data (secret TEXT)')
            live.execute("INSERT INTO unrelated_data VALUES ('not CRE')")
            live.commit()
            before=hashlib.sha256(self.source.read_bytes()).hexdigest()
            report=migrate(self.source,self.target)
            self.assertEqual(before,hashlib.sha256(self.source.read_bytes()).hexdigest())
            self.assertEqual(live.execute('SELECT count(*) FROM admin_users').fetchone()[0],4)
        with closing(sqlite3.connect(str(self.target))) as db:
            tables={row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            self.assertEqual(tables,{'stores','admin_users','customer_product_types','customer_form_fields','cre_staff_members','customer_visits','cre_visit_options','cre_walkout_records','cre_walkout_followups','customer_product_categories','walk_in_feedback','feedback_send_logs'})
            self.assertEqual(db.execute('SELECT customer_name FROM customer_visits').fetchone()[0],'Committed WAL')
            self.assertEqual(db.execute('SELECT id FROM admin_users ORDER BY id').fetchall(),[(1,),(2,),(4,)])
            self.assertEqual(db.execute('PRAGMA foreign_key_check').fetchall(),[])
            self.assertEqual(db.execute('SELECT custom_fields FROM customer_visits').fetchone()[0],'{"1": "5000"}')
        self.assertEqual(report['tables']['customer_visits'],1)
        self.assertEqual(report['integrity'],'ok')

    def test_existing_target_and_source_are_never_overwritten(self):
        migrate(self.source,self.target)
        before=self.target.read_bytes()
        with self.assertRaises(FileExistsError): migrate(self.source,self.target)
        self.assertEqual(before,self.target.read_bytes())
        with self.assertRaises(ValueError): migrate(self.source,self.source)

    def test_broken_relationship_rolls_back_without_publishing(self):
        with closing(sqlite3.connect(str(self.source))) as db, db:
            db.execute('UPDATE customer_visits SET handled_by_id=999 WHERE id=1')
        with self.assertRaises(sqlite3.IntegrityError): migrate(self.source,self.target)
        self.assertFalse(self.target.exists())
        self.assertFalse(list(self.target.parent.glob('.cre-import-*')))

    def test_unknown_schema_is_rejected_without_data_loss(self):
        with closing(sqlite3.connect(str(self.source))) as db, db:
            db.execute('ALTER TABLE customer_visits ADD COLUMN future_value TEXT')
        with self.assertRaises(ValueError): migrate(self.source,self.target)
        self.assertFalse(self.target.exists())

    def test_store_database_migration_preserves_rows_and_applies_assignments(self):
        with closing(sqlite3.connect(str(self.source))) as db, db:
            db.execute("INSERT INTO stores (id,store_name,city,store_address,timings,created_at) VALUES (2,'Second Store','Ujjain','Address','10-7',CURRENT_TIMESTAMP)")
            db.execute('UPDATE admin_users SET store_id=NULL WHERE id=4')
            db.execute('UPDATE cre_staff_members SET store_id=NULL')
            db.execute("UPDATE customer_visits SET status='Billed',store_id=NULL")
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        report = migrate_store_database(
            self.source, self.target, 1,
            user_assignments={'executive': 2}, staff_assignments={'1': 2},
        )
        self.assertEqual(before, hashlib.sha256(self.source.read_bytes()).hexdigest())
        with closing(sqlite3.connect(str(self.target))) as db:
            self.assertEqual(db.execute('SELECT store_id FROM customer_visits WHERE id=1').fetchone(), (1,))
            self.assertEqual(db.execute('SELECT store_id FROM admin_users WHERE id=4').fetchone(), (2,))
            self.assertEqual(db.execute('SELECT id,store_id FROM cre_staff_members ORDER BY id').fetchall(), [(1,2),(2,1)])
            self.assertEqual(db.execute('PRAGMA foreign_key_check').fetchall(), [])
        self.assertEqual(report['integrity'], 'ok')
        self.assertTrue(self.target.with_suffix('.store-migration.json').exists())
        with self.assertRaises(FileExistsError):
            migrate_store_database(self.source, self.target, 1)

class DashboardSummaryTests(unittest.TestCase):
    def test_ist_display_converts_utc_and_aware_timestamps(self):
        from datetime_utils import format_ist, INDIA_TIME
        expected='22 Oct 2026, 12:30 AM IST'
        self.assertEqual(format_ist(datetime(2026,10,21,19)),expected)
        self.assertEqual(format_ist(datetime(2026,10,21,19,tzinfo=timezone.utc)),expected)
        self.assertEqual(format_ist(datetime(2026,10,22,0,30,tzinfo=INDIA_TIME)),expected)
        self.assertEqual(format_ist(datetime(2026,10,21,19),'%d %b %Y'),'22 Oct 2026')
        self.assertEqual(format_ist(None),'-')
    def test_category_migration_preserves_legacy_visits_and_locks_billed(self):
        from visit_categories import migrate_003_categories_and_billed_lock
        engine = create_engine('sqlite:///:memory:')
        try:
            with engine.begin() as db:
                db.exec_driver_sql('CREATE TABLE customer_visits (id INTEGER PRIMARY KEY, status TEXT)')
                db.exec_driver_sql('CREATE TABLE customer_product_categories (id INTEGER PRIMARY KEY)')
                db.exec_driver_sql('CREATE TABLE cre_walkout_followups (id INTEGER PRIMARY KEY, visit_id INTEGER)')
                db.exec_driver_sql("INSERT INTO customer_visits VALUES (1,'Billed'),(2,'Draft')")
            migrate_003_categories_and_billed_lock(engine)
            migrate_003_categories_and_billed_lock(engine)
            with engine.begin() as db:
                self.assertEqual(db.exec_driver_sql('SELECT id,status,category_id FROM customer_visits ORDER BY id').all(),[(1,'Billed',None),(2,'Draft',None)])
                db.exec_driver_sql("UPDATE customer_visits SET status='Billed' WHERE id=2")
            with engine.begin() as db:
                from sqlalchemy.exc import IntegrityError
                with self.assertRaises(IntegrityError):
                    db.exec_driver_sql("UPDATE customer_visits SET status='Draft' WHERE id=2")
        finally:
            engine.dispose()

    def test_walkout_backfill_is_idempotent_and_keeps_original_snapshot(self):
        from walkout_followups import migrate_002_walkout_followups
        from sqlalchemy import text
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'backfill.db'
            fixture(path)
            engine = create_engine('sqlite:///' + path.as_posix())
            try:
                with engine.begin() as db:
                    db.execute(text("UPDATE customer_visits SET status='Walkout', walkout_reason='Budget', remark='Original'"))
                migrate_002_walkout_followups(engine)
                with engine.begin() as db:
                    db.execute(text("UPDATE customer_visits SET status='Billed', walkout_reason=NULL, remark=NULL"))
                migrate_002_walkout_followups(engine)
                with engine.connect() as db:
                    self.assertEqual(db.execute(text('SELECT visit_id,reason,remark FROM cre_walkout_records')).all(), [(1,'Budget','Original')])
                    self.assertEqual(len(db.execute(text("SELECT name FROM sqlite_master WHERE type='trigger' AND name LIKE 'cre_followup_no_%'")).all()), 2)
            finally:
                engine.dispose()

    def test_india_midnight_boundaries_and_daily_reset_preserve_history(self):
        from routers.customer_relations import dashboard_day_summary, INDIA_TIME
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'summary.db'
            fixture(path)
            engine = create_engine('sqlite:///' + path.as_posix())
            try:
                with Session(engine) as db:
                    db.query(CustomerVisit).delete()
                    for index, (stamp, status, activity) in enumerate([
                        ('2026-10-21 18:29:59', 'Billed', 'New Jewellery'),
                        ('2026-10-21 18:30:00', 'Billed', 'New Jewellery'),
                        ('2026-10-22 10:00:00', 'Walkout', 'New Jewellery'),
                        ('2026-10-22 11:00:00', None, 'New Jewellery'),
                        ('2026-10-22 12:00:00', None, 'Old Gold'),
                        ('2026-10-22 18:29:59', 'Draft', 'New Jewellery'),
                        ('2026-10-22 18:30:00', 'Billed', 'New Jewellery'),
                    ]):
                        db.add(CustomerVisit(customer_name='Boundary '+str(index), mobile_number='9876543210',
                            product_type_id=1, created_by=1, activity_type=activity, status=status,
                            created_at=datetime.fromisoformat(stamp)))
                    db.commit()
                    today = dashboard_day_summary(db, datetime(2026, 10, 22, 23, 59, 59, tzinfo=INDIA_TIME))
                    self.assertEqual([today[k] for k in ('total_visits', 'draft_count', 'billed_count', 'walkout_count')], [5, 2, 1, 1])
                    self.assertEqual(today['resets_at'], '2026-10-23T00:00:00+05:30')
                    tomorrow = dashboard_day_summary(db, datetime(2026, 10, 23, tzinfo=INDIA_TIME))
                    self.assertEqual([tomorrow[k] for k in ('total_visits', 'draft_count', 'billed_count', 'walkout_count')], [1, 0, 1, 0])
                    self.assertEqual(db.query(CustomerVisit).count(), 7)
            finally:
                engine.dispose()

class FeedbackWhatsappTests(unittest.TestCase):
    def test_kit19_payload_uses_tokenized_url_and_normalized_mobile(self):
        from services.feedback_whatsapp import send_feedback_whatsapp

        captured = {}
        class Response:
            def __enter__(self): return self
            def __exit__(self, *_): return False
            def read(self): return b'{"meta":{"success":true,"http_code":200,"developer_message":"Message with request id request-1 has been submitted to processing queue successfully."},"success":true}'
        def fake_urlopen(request, timeout):
            captured['url'] = request.full_url
            captured['payload'] = json.loads(request.data)
            captured['timeout'] = timeout
            return Response()

        environment = {
            'KIT19_WHATSAPP_URL': 'https://services.kit19.com/IMS/Whatsapp/Template',
            'KIT19_WHATSAPP_KEY': 'secret-key',
            'KIT19_WHATSAPP_USERNAME': 'anand-user',
            'FEEDBACK_PUBLIC_URL': 'https://anandjewels.com/feedback.html',
        }
        with patch.dict(os.environ, environment), patch(
            'services.feedback_whatsapp.urlopen', side_effect=fake_urlopen
        ):
            result = send_feedback_whatsapp('95998 81068', 'feedback-token')

        payload = captured['payload']
        self.assertEqual(captured['url'], environment['KIT19_WHATSAPP_URL'])
        self.assertEqual(captured['timeout'], 15)
        self.assertEqual(payload['whatsapp']['to'], '919599881068')
        self.assertEqual(payload['whatsapp']['template']['name'], 'feedback_template')
        self.assertEqual(
            payload['whatsapp']['template']['components'][0]['parameters'][0]['text'],
            'https://anandjewels.com/feedback.html?token=feedback-token',
        )
        self.assertEqual(result['provider_request_id'], 'request-1')


class WebTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory()
        cls.dbpath=Path(cls.tmp.name)/'web.db'
        fixture(cls.dbpath)
        with closing(sqlite3.connect(str(cls.dbpath))) as db, db:
            password_hash = db.execute('SELECT password_hash FROM admin_users WHERE id=1').fetchone()[0]
            db.execute("INSERT INTO stores (id,store_name,city,store_address,timings,created_at) VALUES (2,'Second Store','Ujjain','Second address','10-7',CURRENT_TIMESTAMP)")
            db.execute("INSERT INTO admin_users (id,username,display_name,password_hash,role,store_id,created_at) VALUES (5,'second_executive','Second CRE',?,'customer_relation_executive',2,CURRENT_TIMESTAMP)", (password_hash,))
            db.execute("INSERT INTO cre_staff_members (id,name,staff_type,store_id,created_by,created_at,is_active) VALUES (3,'Second Sales','sales_person',2,1,CURRENT_TIMESTAMP,1)")
            db.execute("INSERT INTO cre_staff_members (id,name,staff_type,store_id,created_by,created_at,is_active) VALUES (4,'Second Manager','manager',2,1,CURRENT_TIMESTAMP,1)")
        with socket.socket() as sock:
            sock.bind(('127.0.0.1',0))
            cls.port=sock.getsockname()[1]
        cls.url='http://127.0.0.1:'+str(cls.port)
        env=dict(os.environ, CRE_DATABASE_URL='sqlite:///'+cls.dbpath.as_posix(),
                 CRE_SECRET_KEY='independent-test-secret-that-is-not-used-outside-tests',
                 ANAND_INTEGRATION_KEY='test-integration-key', FEEDBACK_AUTO_SEND='false',
                 PYTHONUTF8='1')
        env.pop('PYTHONPATH',None)
        cls.log=open(Path(cls.tmp.name)/'server.log','w+',encoding='utf-8')
        bootstrap = (
            "import sys,threading; sys.path.insert(0,"+repr(str(ROOT))+"); "
            "import auth,time; from pathlib import Path; original_checkpw=auth.bcrypt.checkpw; "
            "exec(\"def slow_checkpw(password, hashed):\\n"
            "    if password == b'concurrency-probe':\\n"
            "        Path('login-started').touch()\\n"
            "        time.sleep(2)\\n"
            "    return original_checkpw(password, hashed)\\n\"); "
            "auth.bcrypt.checkpw=slow_checkpw; "
            "import uvicorn; server=uvicorn.Server(uvicorn.Config('server:app',host='127.0.0.1',port="+str(cls.port)+")); "
            "threading.Thread(target=lambda: (sys.stdin.buffer.read(1),setattr(server,'should_exit',True)),daemon=True).start(); server.run()"
        )
        cls.process=subprocess.Popen([sys.executable,'-B','-c',bootstrap],
                                     cwd=cls.tmp.name,env=env,stdin=subprocess.PIPE,stdout=cls.log,stderr=subprocess.STDOUT)
        try:
            for _ in range(100):
                if cls.process.poll() is not None: raise RuntimeError('Test server exited')
                try:
                    with build_opener().open(cls.url+'/health',timeout=1) as response:
                        if response.status==200: break
                except OSError: time.sleep(.2)
            else: raise RuntimeError('Test server did not start')
        except Exception:
            cls.stop_server()
            cls.log.seek(0); detail=cls.log.read(); cls.log.close(); cls.tmp.cleanup()
            raise RuntimeError(detail)
    @classmethod
    def stop_server(cls):
        cls.process.communicate(input=b'x',timeout=10)
    @classmethod
    def tearDownClass(cls):
        cls.stop_server()
        cls.log.close(); cls.tmp.cleanup()
    def setUp(self):
        self.cookies=CookieJar()
        self.client=build_opener(HTTPCookieProcessor(self.cookies))
    def request(self,path,data=None):
        body=urlencode(data).encode() if data is not None else None
        try:
            with self.client.open(Request(self.url+path,data=body),timeout=10) as response:
                return response.status,response.read(),response.url,response.headers
        except HTTPError as error:
            return error.code,error.read(),error.url,error.headers
    def api(self, path, payload=None, method=None, integration_key=None):
        body = json.dumps(payload).encode() if payload is not None else None
        headers = {'Accept': 'application/json'}
        if payload is not None:
            headers['Content-Type'] = 'application/json'
        if integration_key is not None:
            headers['X-Integration-Key'] = integration_key
        request = Request(self.url + path, data=body, headers=headers, method=method)
        try:
            with build_opener().open(request, timeout=10) as response:
                return response.status, json.loads(response.read())
        except HTTPError as error:
            return error.code, json.loads(error.read())
    def login(self,user='owner',password=PASSWORD):
        return self.request('/cre/login',{'username':user,'password':password})
    def sql(self,query,params=()):
        with closing(sqlite3.connect(str(self.dbpath))) as db:
            return db.execute(query,params).fetchall()
    def form(self,**changes):
        data=dict(customer_name='Standalone Customer',mobile_number='9876543210',product_type_id=1,
                  handled_by_id=1,assigned_manager_id=2,store_id=1,category_id=1,activity_type='New Jewellery',custom_1='12000')
        data.update(changes); return data
    def test_slow_password_check_does_not_block_other_requests(self):
        marker = Path(self.tmp.name) / 'login-started'
        marker.unlink(missing_ok=True)
        def slow_login():
            client = build_opener()
            try:
                client.open(Request(self.url+'/cre/login', data=urlencode({
                    'username': 'owner', 'password': 'concurrency-probe'
                }).encode()), timeout=10)
            except HTTPError as error:
                return error.code
        with ThreadPoolExecutor(max_workers=1) as workers:
            pending = workers.submit(slow_login)
            deadline = time.monotonic() + 5
            while not marker.exists() and time.monotonic() < deadline:
                time.sleep(.01)
            self.assertTrue(marker.exists(), 'Password verification did not start')
            self.assertEqual(self.request('/health')[0], 200)
            self.assertEqual(self.request('/cre/login')[0], 200)
            self.assertFalse(pending.done(), 'Slow password verification blocked other requests')
            self.assertEqual(pending.result(timeout=10), 401)

    def test_compressed_dashboard_preserves_content(self):
        import gzip
        self.login()
        plain = self.request('/cre/dashboard')[1]
        with self.client.open(Request(self.url+'/cre/dashboard', headers={
            'Accept-Encoding': 'gzip'
        }), timeout=10) as response:
            compressed = response.read()
            self.assertEqual(response.headers.get('Content-Encoding'), 'gzip')
        self.assertEqual(gzip.decompress(compressed), plain)
        self.assertLess(len(compressed), len(plain) // 2)

    def test_concurrent_dashboard_and_report_reads(self):
        self.login()
        cookie = '; '.join(f'{item.name}={item.value}' for item in self.cookies)
        paths = ['/cre/dashboard', '/cre/reports', '/cre/dashboard/summary', '/health'] * 6
        def read(path):
            with build_opener().open(Request(self.url+path, headers={'Cookie': cookie}), timeout=10) as response:
                return response.status, response.url, response.read()
        with ThreadPoolExecutor(max_workers=8) as workers:
            results = list(workers.map(read, paths))
        for path, (status, url, body) in zip(paths, results):
            self.assertEqual(status, 200)
            self.assertEqual(urlparse(url).path, path)
            self.assertTrue(body)

    def test_login_roles_cookie_and_logout(self):
        self.assertEqual(urlparse(self.request('/cre/dashboard')[2]).path,'/cre/login')
    def test_feedback_submission_store_scope_and_incremental_integration(self):
        public_form = self.request('/feedback.html?utm_src=cre&token=local-test-token')[1]
        self.assertIn(b'Customer Feedback | Anand Jewels', public_form)
        self.assertIn(b"fetch('/api/feedback'", public_form)
        self.assertIn(b'local-test-token', self.request('/feedback.html?token=local-test-token')[2].encode())
        self.login()
        first_name = 'Feedback first store'
        second_name = 'Feedback second store'
        self.assertNotIn('error=', self.request('/cre/customers', self.form(
            customer_name=first_name, status='Draft'
        ))[2])
        self.assertNotIn('error=', self.request('/cre/customers', self.form(
            customer_name=second_name, status='Draft', store_id=2,
            handled_by_id=3, assigned_manager_id=4,
        ))[2])
        feedback_rows = self.sql(
            "SELECT walk_in_feedback.id,token,send_status,customer_visits.customer_name "
            "FROM walk_in_feedback JOIN customer_visits ON customer_visits.id=walk_in_feedback.customer_visit_id "
            "WHERE customer_visits.customer_name IN (?,?) ORDER BY walk_in_feedback.id",
            (first_name, second_name),
        )
        self.assertEqual(len(feedback_rows), 2)
        self.assertEqual(len({row[1] for row in feedback_rows}), 2)
        self.assertTrue(all(len(row[1]) >= 40 and row[2] == 'pending' for row in feedback_rows))
        first_id, token = feedback_rows[0][:2]

        self.assertEqual(self.api('/api/feedback', {'token': token, 'rating': 0})[0], 422)
        self.assertEqual(self.api('/api/feedback', {'token': token, 'rating': 6})[0], 422)
        self.assertEqual(self.api('/api/feedback', {'token': 'missing-token-value', 'rating': 5})[0], 404)
        self.assertEqual(self.api('/api/integrations/feedback?limit=500')[0], 401)
        self.assertEqual(self.api('/api/integrations/feedback?limit=500', integration_key='wrong')[0], 401)
        status, initial_sync = self.api(
            '/api/integrations/feedback?limit=500', integration_key='test-integration-key'
        )
        self.assertEqual(status, 200)
        self.assertEqual(len([row for row in initial_sync['items'] if row['id'] == first_id]), 1)
        cursor = initial_sync['next_cursor']

        status, result = self.api('/api/feedback', {
            'token': token, 'rating': 5, 'comments': ' Excellent service '
        })
        self.assertEqual(status, 201)
        self.assertEqual(result['feedback_id'], first_id)
        self.assertEqual(self.api('/api/feedback', {'token': token, 'rating': 4})[0], 409)
        self.assertEqual(self.sql(
            'SELECT rating,comments,send_status,is_read FROM walk_in_feedback WHERE id=?', (first_id,)
        ), [(5, 'Excellent service', 'submitted', 0)])

        query = urlencode({'after_time': cursor['after_time'], 'after_id': cursor['after_id'], 'limit': 500})
        status, changed = self.api(
            '/api/integrations/feedback?' + query, integration_key='test-integration-key'
        )
        self.assertEqual(status, 200)
        self.assertIn(first_id, [row['id'] for row in changed['items']])
        self.assertEqual(self.api(
            f'/api/integrations/feedback/{first_id}/resend', method='POST',
            integration_key='test-integration-key'
        )[0], 409)
        self.assertEqual(self.api(
            f'/api/integrations/feedback/{first_id}/read', method='PATCH',
            integration_key='test-integration-key'
        )[0], 200)

        self.request('/cre/logout')
        self.login(user='executive')
        feedback_page = self.request('/cre/feedback')[1]
        self.assertIn(first_name.encode(), feedback_page)
        self.assertNotIn(second_name.encode(), feedback_page)
        self.assertIn(b'CRE Feedback', feedback_page)
        self.assertEqual(self.login(password='wrong')[0],401)
        self.assertEqual(self.login(user='historical')[0],401)
        self.assertEqual(self.login(user='executive')[0],200)
        self.assertIn('anand_cre_session',{cookie.name for cookie in self.cookies})
        self.assertNotIn('session',{cookie.name for cookie in self.cookies})
        self.assertEqual(urlparse(self.request('/cre/logout')[2]).path,'/cre/login')
        self.assertEqual(urlparse(self.request('/cre/dashboard')[2]).path,'/cre/login')
    def test_optional_sales_person_and_manager(self):
        self.login()
        for index, (sales, manager) in enumerate([(0, 0), (1, 0), (0, 2), ('', '')]):
            data = self.form(customer_name='Optional staff '+str(index), status='Draft', handled_by_id=sales, assigned_manager_id=manager)
            self.assertNotIn('error=', self.request('/cre/customers', data)[2])
            visit, saved_sales, saved_manager = self.sql('SELECT id,handled_by_id,assigned_manager_id FROM customer_visits WHERE customer_name=?', (data['customer_name'],))[0]
            self.assertEqual((saved_sales, saved_manager), (sales or None, manager or None))
            self.assertNotIn('error=', self.request('/cre/customers/'+str(visit)+'/edit', dict(data, handled_by_id=1, assigned_manager_id=2))[2])
            self.assertEqual(self.sql('SELECT handled_by_id,assigned_manager_id FROM customer_visits WHERE id=?', (visit,)), [(1, 2)])
            self.assertNotIn('error=', self.request('/cre/customers/'+str(visit)+'/edit', dict(data, handled_by_id=0, assigned_manager_id=0))[2])
            self.assertEqual(self.sql('SELECT handled_by_id,assigned_manager_id FROM customer_visits WHERE id=?', (visit,)), [(None, None)])
            self.assertIn('error=', self.request('/cre/customers/'+str(visit)+'/edit', dict(data, assigned_manager_id=1))[2])
        self.assertIn('error=', self.request('/cre/customers', self.form(handled_by_id=2))[2])
        self.assertIn('error=', self.request('/cre/customers', self.form(assigned_manager_id=999))[2])
        self.assertIn(b'Not Assigned', self.request('/cre/dashboard')[1])

        self.assertIn(b'Not Assigned', self.request('/cre/reports')[1])
        optional_product = self.form(customer_name='Optional product', product_type_id='', status='Draft')
        self.assertNotIn('error=', self.request('/cre/customers', optional_product)[2])
        self.assertEqual(self.sql(
            "SELECT customer_product_types.name FROM customer_visits "
            "JOIN customer_product_types ON customer_product_types.id=customer_visits.product_type_id "
            "WHERE customer_visits.customer_name='Optional product'"
        ), [('Not Selected',)])
        engine = create_engine('sqlite:///'+self.dbpath.as_posix())
        from sqlalchemy import text
        try:
            with engine.begin() as db:
                db.execute(text('UPDATE cre_staff_members SET is_active=0'))
            self.assertIn(b'id="addCustomerModal"', self.request('/cre/dashboard')[1])
            self.assertNotIn('error=', self.request('/cre/customers', self.form(handled_by_id=0, assigned_manager_id=0))[2])
        finally:
            with engine.begin() as db:
                db.execute(text('UPDATE cre_staff_members SET is_active=1'))
            engine.dispose()

    def test_dashboard_summary_and_newest_first_when_timestamps_match(self):
        self.login()
        for name in ('First tied visit', 'Second tied visit'):
            self.request('/cre/customers', self.form(customer_name=name, status='Draft'))
        with closing(sqlite3.connect(str(self.dbpath))) as db, db:
            db.execute("UPDATE customer_visits SET created_at='2099-01-01 00:00:00' WHERE customer_name IN (?, ?)",
                       ('First tied visit', 'Second tied visit'))
        body = self.request('/cre/dashboard')[1]
        self.assertLess(body.index(b'Second tied visit'), body.index(b'First tied visit'))
        code, data, _, headers = self.request('/cre/dashboard/summary')
        self.assertEqual(code, 200)
        self.assertEqual(headers['Cache-Control'], 'no-store')
        self.assertIn('resets_at', json.loads(data))
        self.request('/cre/logout')
        self.assertEqual(urlparse(self.request('/cre/dashboard/summary')[2]).path, '/cre/login')

    def test_two_users_save_unique_permanent_visit_numbers_at_once(self):
        clients = [WebTests(), WebTests()]
        for client, username in zip(clients, ('owner', 'executive')):
            client.setUp()
            client.login(user=username)
        gate = Barrier(2)
        def save(index):
            gate.wait(timeout=10)
            return clients[index].request('/cre/customers', self.form(
                customer_name='Concurrent user '+str(index), status='Draft'))
        with ThreadPoolExecutor(max_workers=2) as workers:
            results = list(workers.map(save, (0, 1)))
        for result in results:
            self.assertEqual(result[0], 200)
            self.assertNotIn('error=', result[2])
            self.assertIn(b'saved successfully', result[1])
        saved = self.sql("SELECT customer_visits.id,customer_name,created_by,store_visit_number,visit_prefix FROM customer_visits JOIN stores ON stores.id=customer_visits.store_id WHERE customer_name LIKE 'Concurrent user %' ORDER BY customer_visits.id")
        self.assertEqual(len(saved),2)
        self.assertEqual(len({row[3] for row in saved}),2)
        self.assertEqual({row[2] for row in saved},{1,4})
        self.login()
        for visit_id, name, _, sequence, prefix in saved:
            number = f'{prefix}-{sequence:06d}'
            body = self.request('/cre/dashboard?q='+number.lower())[1]
            self.assertIn(number.encode(),body)
            self.assertIn(name.encode(),body)
            other = next(row[1] for row in saved if row[0]!=visit_id)
            self.assertNotIn(other.encode(),body)
            self.assertNotIn('error=',self.request('/cre/customers/'+str(visit_id)+'/edit',self.form(customer_name=name+' edited',status='Draft'))[2])
            self.assertIn(number.encode(),self.request('/cre/dashboard?q='+number)[1])
        self.assertEqual(self.request('/cre/dashboard?q=CRE-9999999999999999999')[0],200)

    def test_single_status_replaces_purpose(self):
        self.login()
        choices = ['Draft', 'Billed', 'Walkout', 'Old Gold', 'Repair Taken', 'Repair Delivered', 'Ordered Delivered']
        for index, choice in enumerate(choices):
            data = self.form(customer_name='Unified status '+str(index), status=choice)
            data.pop('activity_type')
            if choice == 'Walkout':
                data.update(remark='Selected another design', walkout_reason='Design', competitor_shop='Tanishq')
            self.assertNotIn('error=', self.request('/cre/customers', data)[2], choice)
            visit, activity, stored = self.sql('SELECT id,activity_type,status FROM customer_visits WHERE customer_name=?', (data['customer_name'],))[0]
            self.assertEqual((activity, stored), ('New Jewellery', choice) if choice in choices[:3] else (choice, None))
            dashboard = self.request('/cre/dashboard?activity='+quote(choice))[1]
            self.assertIn(data['customer_name'].encode(), dashboard)
            self.assertNotIn(b'<th>Purpose</th>', dashboard)
            self.assertNotIn(b'name="activity_type"', dashboard)
            edit = dict(data, status='Draft')
            if choice == 'Billed':
                self.assertIn('error=', self.request('/cre/customers/'+str(visit)+'/edit', edit)[2])
                continue
            self.assertNotIn('error=', self.request('/cre/customers/'+str(visit)+'/edit', edit)[2])
            self.assertEqual(self.sql('SELECT status,walkout_reason,competitor_shop FROM customer_visits WHERE id=?', (visit,)), [('Draft', None, None)])
            self.assertNotIn('error=', self.request('/cre/customers/'+str(visit)+'/status', {'status': 'Billed'})[2])
        invalid = self.form(status='Not a status'); invalid.pop('activity_type')
        self.assertIn('error=', self.request('/cre/customers', invalid)[2])
        invalid['status'] = 'Walkout'
        self.assertIn('error=', self.request('/cre/customers', invalid)[2])
        self.assertNotIn(b'<th>Purpose</th>', self.request('/cre/reports')[1])

    def test_old_gold_and_repair_statuses_hide_and_clear_product_fields(self):
        self.login()
        hidden_statuses = ('Old Gold', 'Repair Taken', 'Repair Delivered')
        not_selected_id = self.sql(
            "SELECT id FROM customer_product_types WHERE name='Not Selected'"
        )[0][0]
        for index, status in enumerate(hidden_statuses):
            name = 'No product status ' + str(index)
            self.assertNotIn('error=', self.request(
                '/cre/customers', self.form(
                    customer_name=name, status=status, product_type_id='', category_id=''
                )
            )[2])
            visit_id, product_type_id, category_id = self.sql(
                'SELECT id,product_type_id,category_id FROM customer_visits WHERE customer_name=?',
                (name,),
            )[0]
            self.assertEqual((product_type_id, category_id), (not_selected_id, None))
            html = self.request('/cre/dashboard?q=' + quote(name))[1]
            self.assertNotIn(b'<dt>Product Type</dt>', html)
            self.assertNotIn(b'<dt>Category Type</dt>', html)

            draft_name = name + ' draft'
            self.assertNotIn('error=', self.request(
                '/cre/customers', self.form(customer_name=draft_name, status='Draft')
            )[2])
            draft_id = self.sql(
                'SELECT id FROM customer_visits WHERE customer_name=?', (draft_name,)
            )[0][0]
            self.assertNotIn('error=', self.request(
                '/cre/customers/' + str(draft_id) + '/status', {'status': status}
            )[2])
            self.assertEqual(self.sql(
                'SELECT product_type_id,category_id FROM customer_visits WHERE id=?', (draft_id,)
            ), [(not_selected_id, None)])

    def test_dynamic_walkout_choices_and_history(self):
        self.assertEqual(urlparse(self.request('/cre/walkout-options')[2]).path, '/cre/login')
        self.request('/cre/walkout-options/reason', {'name': 'Unauthenticated'})
        self.assertFalse(self.sql("SELECT id FROM cre_visit_options WHERE name='Unauthenticated'"))
        self.login()
        for kind, name in [('reason', 'Delivery timeline'), ('competitor', 'New Local Jewels')]:
            self.assertNotIn('error=', self.request('/cre/walkout-options/'+kind, {'name': name})[2])
            self.assertIn(name.encode(), self.request('/cre/dashboard')[1])
            self.assertIn('error=', self.request('/cre/walkout-options/'+kind, {'name': name.upper()})[2])
        self.assertIn('error=', self.request('/cre/walkout-options/reason', {'name': 'x'*81})[2])
        self.assertIn('error=', self.request('/cre/walkout-options/reason', {'name': '  '})[2])
        self.request('/cre/customers', self.form(customer_name='Dynamic visit'))
        visit = self.sql("SELECT id FROM customer_visits WHERE customer_name='Dynamic visit'")[0][0]
        route = '/cre/customers/'+str(visit)
        payload = dict(status='Walkout', remark='Requested later delivery', walkout_reason='Delivery timeline', competitor_shop='New Local Jewels')
        self.assertNotIn('error=', self.request(route+'/status', payload)[2])
        self.assertNotIn('error=', self.request(route+'/edit', self.form(customer_name='Dynamic visit', **payload))[2])
        self.assertIn('error=', self.request(route+'/status', dict(payload, competitor_shop='Unknown shop'))[2])
        ids = self.sql("SELECT id FROM cre_visit_options WHERE name IN ('Delivery timeline','New Local Jewels')")
        for (option_id,) in ids:
            self.request('/cre/walkout-options/'+str(option_id)+'/remove', {})
        # Migration is repeatable and does not re-enable removed choices.
        from visit_options import migrate_001_visit_options
        engine = create_engine('sqlite:///'+self.dbpath.as_posix())
        try:
            migrate_001_visit_options(engine)
        finally:
            engine.dispose()
        self.assertEqual(self.sql("SELECT sum(is_active) FROM cre_visit_options WHERE name IN ('Delivery timeline','New Local Jewels')"), [(0,)])
        self.assertIn(b'Delivery timeline (removed)', self.request('/cre/dashboard')[1])
        self.assertNotIn('error=', self.request(route+'/edit', self.form(customer_name='Dynamic visit', **payload))[2])
        self.assertIn(b'Delivery timeline', self.request('/cre/reports')[1])
        self.request('/cre/customers', self.form(customer_name='Another dynamic visit'))
        another = self.sql("SELECT id FROM customer_visits WHERE customer_name='Another dynamic visit'")[0][0]
        self.assertIn('error=', self.request('/cre/customers/'+str(another)+'/status', payload)[2])
        for (option_id,) in ids:
            self.request('/cre/walkout-options/'+str(option_id)+'/restore', {})
        self.assertNotIn('error=', self.request('/cre/customers/'+str(another)+'/status', payload)[2])
        self.assertIn(b'Walkout &amp; Competitors', self.request('/cre/walkout-options')[1])
    def test_pages_and_main_admin_are_isolated(self):
        self.login()
        for path in ('/cre/dashboard','/cre/reports','/cre/sales-people','/cre/managers','/cre/products','/cre/fields','/cre/walkout-options'):
            code,body,_,_=self.request(path)
            self.assertEqual(code,200,path)
            self.assertNotIn(b'href="/admin',body)
            self.assertNotIn(b'Gold Rate Management',body)
        self.assertEqual(self.request('/static/images/logo.png')[0],200)
        self.assertEqual(self.request('/static/style.css')[0],200)
        self.assertEqual(self.request('/admin')[0],404)
        self.assertEqual(self.request('/api/gold-rates/latest')[0],404)
    def test_customer_create_edit_walkout_and_export(self):
        self.login()
        _,_,url,_=self.request('/cre/customers',self.form())
        self.assertNotIn('error=',url)
        visit=self.sql("SELECT id FROM customer_visits WHERE customer_name='Standalone Customer' ORDER BY id DESC")[0][0]
        self.assertEqual(self.sql('SELECT status,created_by,custom_fields FROM customer_visits WHERE id=?',(visit,)),[('Draft',1,'{"1": "12000"}')])
        self.request('/cre/customers/'+str(visit)+'/edit',self.form(customer_name='Edited Customer',status='Draft'))
        self.assertEqual(self.sql('SELECT customer_name FROM customer_visits WHERE id=?',(visit,))[0][0],'Edited Customer')
        self.request('/cre/customers/'+str(visit)+'/status',dict(status='Walkout',remark='Budget test',walkout_reason='Budget',competitor_shop='Tanishq'))
        self.assertEqual(self.sql('SELECT status,walkout_reason FROM customer_visits WHERE id=?',(visit,)),[('Walkout','Budget')])
        filters='?date_from=2000-01-01&date_to=2100-01-01&store_id=1'
        code,body,_,_=self.request('/cre/reports'+filters)
        self.assertEqual(code,200)
        self.assertIn(b'Budget',body)
        code,body,_,headers=self.request('/cre/reports/export/excel'+filters)
        self.assertEqual(code,200)
        workbook=load_workbook(io.BytesIO(body))
        self.assertEqual(len(workbook.sheetnames),6)
        self.assertGreaterEqual(workbook['Walk-in Summary']['B5'].value,1)
        self.assertIn('.xlsx',headers['Content-Disposition'])
        self.assertIn(b'Edited Customer',self.request('/cre/dashboard?q=Edited')[1])

    def test_each_report_tab_downloads_its_filtered_excel(self):
        self.login()
        self.request('/cre/customers', self.form(
            customer_name='Excel section customer', status='Walkout',
            remark='Excel section remark', walkout_reason='Budget', competitor_shop='Tanishq',
        ))
        filters = '?date_from=2000-01-01&date_to=2100-01-01&store_id=1'
        expected = {
            'overview': ('Overview', 'Metric'),
            'walkin': ('Walk-in Report', 'Date & time (IST)'),
            'walkout': ('Walkout Report', 'Date & time (IST)'),
            'competitor': ('Competitor Report', 'Date & time (IST)'),
            'manager': ('Manager Report', 'Manager'),
            'sales': ('Sales Person Report', 'Sales Person'),
        }
        page = self.request('/cre/reports' + filters)[1]
        self.assertIn(b'Download Complete Excel', page)
        for section, (sheet_name, first_header) in expected.items():
            self.assertIn(f'section={section}'.encode(), page)
            code, body, _, headers = self.request(
                '/cre/reports/export/excel' + filters + '&section=' + section
            )
            self.assertEqual(code, 200, section)
            self.assertIn('.xlsx', headers['Content-Disposition'])
            workbook = load_workbook(io.BytesIO(body))
            self.assertEqual(workbook.sheetnames, [sheet_name])
            self.assertEqual(workbook[sheet_name]['A1'].value, first_header)
        competitor_book = load_workbook(io.BytesIO(self.request(
            '/cre/reports/export/excel' + filters + '&section=competitor'
        )[1]))
        self.assertTrue(any(
            row[2] == 'Excel section customer'
            for row in competitor_book['Competitor Report'].iter_rows(min_row=2, values_only=True)
        ))
        self.assertEqual(
            self.request('/cre/reports/export/excel' + filters + '&section=invalid')[0], 422
        )

    def test_required_custom_field_and_invalid_relationship_are_rejected(self):
        self.login()
        before=self.sql('SELECT count(*) FROM customer_visits')[0][0]
        self.assertIn('error=',self.request('/cre/customers',self.form(custom_1=''))[2])
        self.assertIn('error=',self.request('/cre/customers',self.form(handled_by_id=999))[2])
        self.assertEqual(self.sql('SELECT count(*) FROM customer_visits')[0][0],before)

    def test_walkout_followups_lock_and_conversion_updates_reports(self):
        self.login()
        payload = self.form(customer_name='Followup customer', status='Walkout',
                            remark='Original walkout remark', walkout_reason='Budget', competitor_shop='Tanishq')
        self.assertNotIn('error=', self.request('/cre/customers', payload)[2])
        visit = self.sql("SELECT id FROM customer_visits WHERE customer_name='Followup customer'")[0][0]
        route = '/cre/reports/walkouts/'+str(visit)
        return_to = '/cre/reports?tab=walkout&date_from=2000-01-01&date_to=2100-01-01&store_id=1'
        self.assertIn('error=', self.request(route+'/followups/2', {'note':'Too early'})[2])
        self.assertIn('error=', self.request(route+'/followups/4', {'note':'Invalid'})[2])
        self.assertIn('error=', self.request(route+'/followups/1', {'note':'  '})[2])
        for slot in (1, 2, 3):
            response = self.request(route+'/followups/'+str(slot), {'note':'Call '+str(slot), 'return_to':return_to})
            self.assertNotIn('error=', response[2])
            self.assertIn('store_id=1', response[2])
            self.assertIn('error=', self.request(route+'/followups/'+str(slot), {'note':'Overwrite'})[2])
        self.assertEqual(self.sql('SELECT slot,note FROM cre_walkout_followups WHERE visit_id=? ORDER BY slot', (visit,)),
                         [(1,'Call 1'), (2,'Call 2'), (3,'Call 3')])
        for statement in ('UPDATE cre_walkout_followups SET note=\'Changed\' WHERE visit_id=?',
                          'DELETE FROM cre_walkout_followups WHERE visit_id=?'):
            with closing(sqlite3.connect(str(self.dbpath))) as db:
                with self.assertRaises(sqlite3.IntegrityError):
                    db.execute(statement, (visit,))
        filters = '?date_from=2000-01-01&date_to=2100-01-01&store_id=1'
        before = load_workbook(io.BytesIO(self.request('/cre/reports/export/excel'+filters)[1]))
        self.assertIn('error=', self.request(route+'/outcome', {'status':'Draft'})[2])
        self.assertNotIn('error=', self.request(route+'/outcome', {'status':'Billed','return_to':return_to})[2])
        after = load_workbook(io.BytesIO(self.request('/cre/reports/export/excel'+filters)[1]))
        self.assertEqual(after['Walk-in Summary']['B3'].value, before['Walk-in Summary']['B3'].value+1)
        self.assertEqual(after['Walk-in Summary']['B5'].value, before['Walk-in Summary']['B5'].value-1)
        register = list(after['Walkout register'].values)
        row = next(row for row in register if row[1]=='Followup customer')
        self.assertEqual(row[10:14], ('Call 1','Call 2','Call 3','Billed'))
        html = self.request('/cre/reports'+filters+'&tab=walkout')[1]
        self.assertIn(b'Followup customer', html)
        self.assertIn(b'Saved &amp; locked', html)
        self.assertIn(b'Original walkout remark', html)
        # Editing final status through the dashboard must preserve register history too.
        self.assertIn('error=', self.request('/cre/customers/'+str(visit)+'/edit', dict(payload,status='Billed',remark='New remark'))[2])
        self.assertEqual(self.sql('SELECT reason,remark FROM cre_walkout_records WHERE visit_id=?', (visit,)),
                         [('Budget','Original walkout remark')])
        self.request('/cre/logout')
        self.assertEqual(urlparse(self.request(route+'/followups/1', {'note':'Unauthenticated'})[2]).path, '/cre/login')

    def test_walkout_snapshot_survives_dashboard_status_conversion(self):
        self.login()
        self.request('/cre/customers', self.form(customer_name='Dashboard recovery',status='Walkout',
                     remark='Keep history',walkout_reason='Budget'))
        visit = self.sql("SELECT id FROM customer_visits WHERE customer_name='Dashboard recovery'")[0][0]
        self.request('/cre/customers/'+str(visit)+'/status', {'status':'Billed'})
        self.assertEqual(self.sql('SELECT status,walkout_reason FROM customer_visits WHERE id=?',(visit,)), [('Billed',None)])
        html = self.request('/cre/reports?tab=walkout&date_from=2000-01-01&date_to=2100-01-01&walkout_reason=Budget')[1]
        self.assertIn(b'Dashboard recovery',html)
        self.assertIn(b'Keep history',html)

    def test_billed_is_locked_across_routes_and_database(self):
        self.login()
        self.request('/cre/customers', self.form(customer_name='Locked billed customer', status='Walkout',
                     walkout_reason='Budget', remark='Original note'))
        visit = self.sql("SELECT id FROM customer_visits WHERE customer_name='Locked billed customer'")[0][0]
        route = '/cre/reports/walkouts/'+str(visit)
        self.assertNotIn('error=', self.request(route+'/outcome', {'status':'Billed'})[2])
        before = self.sql('SELECT * FROM customer_visits WHERE id=?',(visit,))
        for path, payload in [
            ('/cre/customers/'+str(visit)+'/edit', self.form(customer_name='Tampered',status='Draft')),
            ('/cre/customers/'+str(visit)+'/status', {'status':'Draft'}),
            (route+'/outcome', {'status':'Walkout'}),
            (route+'/followups/1', {'note':'After billing'}),
        ]:
            self.assertIn('error=', self.request(path,payload)[2],path)
        self.assertEqual(self.sql('SELECT * FROM customer_visits WHERE id=?',(visit,)),before)
        for statement in ("UPDATE customer_visits SET status='Draft' WHERE id=?", 'DELETE FROM customer_visits WHERE id=?'):
            with closing(sqlite3.connect(str(self.dbpath))) as db:
                with self.assertRaises(sqlite3.IntegrityError):
                    db.execute(statement,(visit,))
        dashboard = self.request('/cre/dashboard?q=Locked+billed')[1]
        self.assertNotIn(('id="editCustomerForm-'+str(visit)+'"').encode(),dashboard)
        self.assertNotIn(b'data-cre-action="edit"',dashboard)
        self.assertIn(b'Billed entry locked',dashboard)

    def test_category_options_optional_and_saved(self):
        self.login()
        self.assertNotIn('error=',self.request('/cre/categories',{'name':'Diamond'})[2])
        self.assertIn('error=',self.request('/cre/categories',{'name':'diamond'})[2])
        self.assertIn('error=',self.request('/cre/categories',{'name':'   '})[2])
        category = self.sql("SELECT id FROM customer_product_categories WHERE name='Diamond'")[0][0]
        for invalid in (0,999999):
            self.assertIn('error=',self.request('/cre/customers',self.form(category_id=invalid))[2])
        self.assertNotIn('error=',self.request('/cre/customers',self.form(
            customer_name='No category customer', status='Draft', category_id=''
        ))[2])
        self.assertEqual(self.sql(
            "SELECT category_id FROM customer_visits WHERE customer_name='No category customer'"
        ), [(None,)])
        self.assertIn(b'Not Selected', self.request('/cre/dashboard?q=No+category+customer')[1])
        self.assertNotIn('error=',self.request('/cre/customers',self.form(customer_name='Category customer',status='Draft',category_id=category))[2])
        visit = self.sql("SELECT id FROM customer_visits WHERE customer_name='Category customer'")[0][0]
        self.assertEqual(self.sql('SELECT category_id FROM customer_visits WHERE id=?',(visit,)),[(category,)])
        self.assertNotIn('error=',self.request('/cre/customers/'+str(visit)+'/edit',self.form(customer_name='Category customer',status='Draft',category_id=1))[2])
        self.assertEqual(self.sql('SELECT category_id FROM customer_visits WHERE id=?',(visit,)),[(1,)])
        html = self.request('/cre/dashboard?q=Category+customer')[1]
        self.assertIn(b'name="category_id"',html)
        self.assertIn(b'id="edit-category-', html)
        self.assertIn(b'data-search-placeholder="Search category..."', html)
        self.assertIn(b'data-search-placeholder="Search product type..."', html)
        self.assertIn(b'Diamond',html)

    def test_dedicated_walkout_register_and_exact_excel_columns(self):
        self.assertEqual(urlparse(self.request('/cre/walkout-register')[2]).path,'/cre/login')
        self.assertEqual(urlparse(self.request('/cre/walkout-register/export/excel')[2]).path,'/cre/login')
        self.login()
        self.request('/cre/customers',self.form(customer_name='Dedicated old walkout',status='Walkout',walkout_reason='Budget',remark='Historic'))
        old = self.sql("SELECT id FROM customer_visits WHERE customer_name='Dedicated old walkout'")[0][0]
        with closing(sqlite3.connect(str(self.dbpath))) as db, db:
            db.execute("UPDATE customer_visits SET created_at='2001-01-01 10:00:00' WHERE id=?",(old,))
        self.request('/cre/customers',self.form(customer_name='Dedicated recovery',status='Walkout',walkout_reason='Budget',remark='Returning'))
        recovered = self.sql("SELECT id FROM customer_visits WHERE customer_name='Dedicated recovery'")[0][0]
        route='/cre/reports/walkouts/'+str(recovered)
        return_to='/cre/walkout-register?store_id=1&manager_id=2'
        response=self.request(route+'/followups/1',{'note':'=2+2','return_to':return_to})
        self.assertEqual(urlparse(response[2]).path,'/cre/walkout-register')
        self.assertIn('manager_id=2',response[2])
        self.request(route+'/outcome',{'status':'Billed','return_to':return_to})
        code, html, _, _ = self.request('/cre/walkout-register')
        self.assertEqual(code,200)
        self.assertIn(b'id="walkout-from" type="date" name="date_from" value="2026-01-01"', html)
        self.assertNotIn(b'Dedicated old walkout', html)
        self.assertIn(b'Dedicated recovery', html)
        default_book=load_workbook(io.BytesIO(self.request('/cre/walkout-register/export/excel')[1]))
        self.assertNotIn('Dedicated old walkout',[row[1] for row in list(default_book.active.values)[1:]])
        legacy = self.request('/cre/walkout-register?date_from=1900-01-01')[1]
        self.assertIn(b'id="walkout-from" type="date" name="date_from" value="2026-01-01"', legacy)
        self.assertNotIn(b'Dedicated old walkout', legacy)
        code, html, _, _ = self.request('/cre/walkout-register?date_from=2001-01-01')
        self.assertEqual(code,200)
        self.assertIn(b'Dedicated old walkout',html)
        self.assertIn(b'Dedicated recovery',html)
        self.assertIn(b'Follow-up pending',html)
        self.assertNotIn(b'id="reportTabs"',html)
        self.assertNotIn(b'id="overview-panel"',html)
        self.assertIn(b'/cre/walkout-register/export/excel?',html)
        code, content, _, headers=self.request('/cre/walkout-register/export/excel')
        self.assertEqual(code,200)
        book=load_workbook(io.BytesIO(content))
        self.assertEqual(book.sheetnames,['Walkout Register'])
        sheet=book.active
        self.assertEqual(list(next(sheet.values)), ['Visit Number','Name','Mobile','Product Type','Category Type','Sales Person','Meet by Manager','Walkout Reason','Status','Follow up 1','Follow up 2','Follow up 3'])
        recovered_row=next(row for row in sheet.iter_rows(min_row=2) if row[1].value=='Dedicated recovery')
        prefix, sequence = self.sql(
            "SELECT stores.visit_prefix,customer_visits.store_visit_number FROM customer_visits "
            "JOIN stores ON stores.id=customer_visits.store_id WHERE customer_visits.id=?", (recovered,)
        )[0]
        self.assertEqual([cell.value for cell in recovered_row[:10]], [f'{prefix}-{sequence:06d}','Dedicated recovery','9876543210','Ring','Gold','Sales Test','Manager Test','Budget','Billed','=2+2'])
        self.assertEqual(recovered_row[9].data_type,'s')
        self.assertEqual(recovered_row[2].data_type,'s')
        self.assertEqual(sheet.freeze_panes,'A2')
        self.assertTrue(sheet.auto_filter.ref)
        self.assertIn('cre-walkout-register.xlsx',headers['Content-Disposition'])
        filtered=load_workbook(io.BytesIO(self.request('/cre/walkout-register/export/excel?date_from=2001-01-01&date_to=2001-01-01')[1]))
        self.assertIn('Dedicated old walkout',[row[1] for row in list(filtered.active.values)[1:]])
        self.assertNotIn('Dedicated recovery',[row[1] for row in list(filtered.active.values)[1:]])

    def test_blank_report_dates_allow_search_and_export(self):
        self.login()
        self.request('/cre/customers', self.form(customer_name='Blank date search',
            status='Walkout', walkout_reason='Budget', remark='Date regression'))
        for path in ('/cre/walkout-register', '/cre/reports'):
            code, html, _, _ = self.request(path + '?date_from=&date_to=&q=Blank+date+search&tab=walkout')
            self.assertEqual(code, 200)
            self.assertIn(b'Blank date search', html)
            self.assertEqual(self.request(path + '/export/excel?date_from=&date_to=&q=Blank+date+search')[0], 200)
            self.assertEqual(self.request(path + '?date_from=2026-02-30')[0], 422)
            self.assertEqual(self.request(path + '/export/excel?date_to=not-a-date')[0], 422)
        self.assertEqual(self.request('/cre/walkout-register?date_from=2001-01-01&date_to=&q=Blank+date+search')[0], 200)

    def test_walkout_search_name_mobile_reference_and_export(self):
        self.login()
        for index in range(12):
            self.request('/cre/customers', self.form(
                customer_name='Search Person ' + str(index), mobile_number='9123456789',
                status='Walkout', walkout_reason='Budget', remark='Search test'))
        ids = [row[0] for row in self.sql(
            "SELECT id FROM customer_visits WHERE customer_name LIKE 'Search Person %' ORDER BY id")]
        self.assertEqual(len(ids), 12)
        self.request('/cre/customers', self.form(customer_name='Unrelated search visitor',
            mobile_number='9988776655', status='Walkout', walkout_reason='Budget', remark='Other'))
        for term, expected in [('search person', 12), ('23456789', 12),
                               (f'cre-{ids[0]:06d}', 1),
                               ('no such customer', 0), ('%', 0), ('_', 0),
                               ('CRE-9999999999999999999', 0)]:
            query = urlencode({'q': term, 'per_page': 10, 'walkout_page': 2})
            code, html, _, _ = self.request('/cre/walkout-register?' + query)
            self.assertEqual(code, 200)
            self.assertIn(b'id="walkout-search"', html)
            self.assertIn(('of ' + str(expected)).encode(), html)
            table = html.split(b'<table class="table table-hover report-walkout-table', 1)[1].split(b'</table>', 1)[0]
            self.assertNotIn(b'Unrelated search visitor', table)
            book = load_workbook(io.BytesIO(self.request('/cre/walkout-register/export/excel?' + query)[1]))
            self.assertEqual(book.active.max_row, expected + 1)
            if expected == 12:
                self.assertEqual(table.count(b'<div class="fw-semibold">Search Person'), 2)
                self.assertIn(b'q=', html)
        redirect = self.request('/cre/reports/walkouts/' + str(ids[0]) + '/followups/1', {
            'note': 'Search follow-up',
            'return_to': '/cre/walkout-register?q=Search+Person&per_page=10&walkout_page=2'})
        self.assertIn('q=Search+Person', redirect[2])
        self.assertIn('walkout_page=2', redirect[2])

    def test_walkout_search_ignores_dates_but_keeps_other_filters(self):
        self.login()
        self.request('/cre/customers', self.form(customer_name='All dates search probe',
            status='Walkout', walkout_reason='Budget', remark='Historical search'))
        visit = self.sql("SELECT id FROM customer_visits WHERE customer_name='All dates search probe'")[0][0]
        with closing(sqlite3.connect(str(self.dbpath))) as db, db:
            db.execute("UPDATE customer_visits SET created_at='2001-01-01 10:00:00' WHERE id=?", (visit,))
        for path in ('/cre/walkout-register', '/cre/reports'):
            for dates in ('date_from=2026-01-01&date_to=2026-01-02', 'date_from=&date_to=', 'date_from=bad&date_to=bad'):
                query = '?tab=walkout&q=All+dates+search+probe&store_id=1&' + dates
                code, html, _, _ = self.request(path + query)
                self.assertEqual(code, 200)
                self.assertIn(b'All dates search probe', html)
                self.assertIn(b'Search results across all dates', html)
                book = load_workbook(io.BytesIO(self.request(path + '/export/excel' + query)[1]))
                sheet = book['Walkout Register' if path.endswith('walkout-register') else 'Walkout register']
                self.assertEqual(sheet.max_row, 2)
            filtered = self.request(path + '?tab=walkout&q=All+dates+search+probe&store_id=999')[1]
            self.assertNotIn(b'<div class="fw-semibold">All dates search probe', filtered)
            without_search = self.request(path + '?tab=walkout&date_from=2026-01-01&date_to=2026-01-02')[1]
            self.assertNotIn(b'<div class="fw-semibold">All dates search probe', without_search)

    def test_register_edit_only_changes_status_reason_and_competitor(self):
        self.login()
        self.request('/cre/customers',self.form(customer_name='Restricted edit customer',status='Walkout',walkout_reason='Budget',remark='Keep remark',competitor_shop='Tanishq'))
        visit=self.sql("SELECT id FROM customer_visits WHERE customer_name='Restricted edit customer'")[0][0]
        route='/cre/reports/walkouts/'+str(visit)
        return_to='/cre/walkout-register?store_id=1'
        self.request(route+'/followups/1',{'note':'Keep follow-up','return_to':return_to})
        locked_columns='customer_name,mobile_number,product_type_id,category_id,handled_by_id,assigned_manager_id,store_id,remark'
        before=self.sql('SELECT '+locked_columns+' FROM customer_visits WHERE id=?',(visit,))
        payload=dict(status='Walkout',walkout_reason='Design',competitor_shop='Caratlane',return_to=return_to)
        self.assertIn('error=',self.request(route+'/edit',dict(payload,customer_name='Tampered'))[2])
        self.assertIn('error=',self.request(route+'/edit',dict(payload,walkout_reason='Unknown'))[2])
        self.assertIn('error=',self.request(route+'/edit',dict(payload,competitor_shop='Unknown'))[2])
        response=self.request(route+'/edit',payload)
        self.assertNotIn('error=',response[2])
        self.assertEqual(urlparse(response[2]).path,'/cre/walkout-register')
        self.assertEqual(self.sql('SELECT '+locked_columns+' FROM customer_visits WHERE id=?',(visit,)),before)
        self.assertEqual(self.sql('SELECT status,walkout_reason,competitor_shop FROM customer_visits WHERE id=?',(visit,)),[('Walkout','Design','Caratlane')])
        self.assertEqual(self.sql('SELECT note FROM cre_walkout_followups WHERE visit_id=?',(visit,)),[('Keep follow-up',)])
        html=self.request('/cre/walkout-register')[1]
        self.assertIn(b'id="walkoutEditModal"',html)
        self.assertIn(b'data-reason="Design"',html)
        self.assertIn(b'data-competitor="Caratlane"',html)
        self.assertNotIn('error=',self.request(route+'/edit',dict(payload,status='Billed'))[2])
        self.assertIn('error=',self.request(route+'/edit',payload)[2])
        # The dashboard status route clears live walkout fields; reports must retain latest choices.
        self.assertEqual(self.sql('SELECT reason,competitor FROM cre_walkout_records WHERE visit_id=?',(visit,)),[('Design','Caratlane')])
        export=load_workbook(io.BytesIO(self.request('/cre/walkout-register/export/excel')[1]))
        exported=next(row for row in list(export.active.values)[1:] if row[1]=='Restricted edit customer')
        self.assertEqual(exported[7:10],('Design','Billed','Keep follow-up'))
        self.request('/cre/logout')
        self.assertEqual(urlparse(self.request(route+'/edit',payload)[2]).path,'/cre/login')

    def test_visit_and_followup_times_are_ist_in_pages_and_excel(self):
        self.login()
        self.request('/cre/customers',self.form(customer_name='IST time customer',status='Walkout',walkout_reason='Budget',remark='Timezone check'))
        visit=self.sql("SELECT id FROM customer_visits WHERE customer_name='IST time customer'")[0][0]
        with closing(sqlite3.connect(str(self.dbpath))) as db,db:
            db.execute("UPDATE customer_visits SET created_at='2026-10-21 19:00:00' WHERE id=?",(visit,))
            db.execute("INSERT INTO cre_walkout_followups(visit_id,slot,note,created_by,created_at) VALUES (?,1,'Evening call',1,'2026-10-21 19:15:00')",(visit,))
        expected=b'22 Oct 2026, 12:30 AM IST'
        self.assertIn(expected,self.request('/cre/dashboard?q=IST+time+customer')[1])
        filters='?date_from=2026-10-22&date_to=2026-10-22'
        report=self.request('/cre/walkout-register'+filters)[1]
        self.assertIn(expected,report)
        self.assertIn(b'22 Oct 2026, 12:45 AM IST',report)
        self.assertNotIn(b' UTC',report)
        book=load_workbook(io.BytesIO(self.request('/cre/reports/export/excel'+filters)[1]))
        row=next(row for row in list(book['Walkout register'].values)[1:] if row[1]=='IST time customer')
        self.assertEqual(row[0],expected.decode())
        self.assertEqual(self.sql('SELECT created_at FROM customer_visits WHERE id=?',(visit,)),[('2026-10-21 19:00:00',)])

    def test_pagination_preserves_filters_totals_and_full_excel(self):
        self.login()
        with closing(sqlite3.connect(str(self.dbpath))) as db,db:
            db.executemany("INSERT INTO customer_visits (customer_name,mobile_number,product_type_id,category_id,store_id,handled_by_id,assigned_manager_id,activity_type,status,walkout_reason,remark,created_by,created_at,updated_at) VALUES (?, '9876543210',1,1,1,1,2,'New Jewellery','Walkout','Budget','Pagination',1,'2005-01-01 10:00:00','2005-01-01 10:00:00')", [('Pagination probe '+str(i),) for i in range(45)])
        ids=[row[0] for row in self.sql("SELECT id FROM customer_visits WHERE customer_name LIKE 'Pagination probe %' ORDER BY id DESC")]
        try:
            for page,first,last in ((1,0,20),(2,20,40),(3,40,45),(999,40,45)):
                code,html,_,_=self.request('/cre/dashboard?q=Pagination+probe&page='+str(page))
                self.assertEqual(code,200)
                table=html.split(b'<table class="table cre-table',1)[1].split(b'</table>',1)[0]
                for index,visit in enumerate(ids):
                    self.assertEqual(f'CRE-{visit:06d}'.encode() in table,first<=index<last)
                self.assertIn(b'of 45',html)
                self.assertIn(b'Previous',html)
                self.assertIn(b'Next',html)
            filters='date_from=2005-01-01&date_to=2005-01-01&store_id=1&manager_id=2'
            for path,param in (('/cre/reports?tab=walkin&','walkin_page'),('/cre/walkout-register?','walkout_page')):
                code,html,_,_=self.request(path+filters+'&'+param+'=2')
                self.assertEqual(code,200)
                self.assertIn(b'of 45',html)
                self.assertIn(b'store_id=1',html)
                self.assertIn(b'manager_id=2',html)
                if param=='walkout_page':
                    table=html.split(b'<table class="table table-hover report-walkout-table',1)[1].split(b'</table>',1)[0]
                else:
                    table=html.split(b'Walk-in register',1)[1].split(b'<table',1)[1].split(b'</table>',1)[0]
                for index,visit in enumerate(ids):
                    self.assertEqual(f'CRE-{visit:06d}'.encode() in table,20<=index<40)
            exported=load_workbook(io.BytesIO(self.request('/cre/walkout-register/export/excel?'+filters+'&walkout_page=2')[1]))
            self.assertEqual(exported.active.max_row,46)
            self.assertEqual({row[0] for row in list(exported.active.values)[1:]},{f'CRE-{visit:06d}' for visit in ids})
            for size, count in ((10,10),(50,45),(100,45)):
                for path,marker in (
                    ('/cre/dashboard?q=Pagination+probe&','<table class="table cre-table'),
                    ('/cre/reports?tab=walkin&'+filters+'&','Walk-in register'),
                    ('/cre/walkout-register?'+filters+'&','<table class="table table-hover report-walkout-table'),
                ):
                    code,html,_,_=self.request(path+'per_page='+str(size))
                    self.assertEqual(code,200)
                    table=html.split(marker.encode(),1)[1].split(b'</table>',1)[0]
                    self.assertEqual(sum(f'CRE-{visit:06d}'.encode() in table for visit in ids),count)
                    self.assertIn(('value="'+str(size)+'" selected').encode(),html)
                    self.assertIn(b'Entries per page',html)
            second=self.request('/cre/walkout-register?'+filters+'&per_page=10&walkout_page=2')[1]
            table=second.split(b'<table class="table table-hover report-walkout-table',1)[1].split(b'</table>',1)[0]
            for index,visit in enumerate(ids):
                self.assertEqual(f'CRE-{visit:06d}'.encode() in table,10<=index<20)
            self.assertIn(b'per_page=10',second)
            full=load_workbook(io.BytesIO(self.request('/cre/walkout-register/export/excel?'+filters+'&per_page=10')[1]))
            self.assertEqual(full.active.max_row,46)
            self.assertEqual(self.request('/cre/dashboard?per_page=1000')[0],422)
            target='/cre/walkout-register?'+filters+'&walkout_page=2'
            response=self.request('/cre/reports/walkouts/'+str(ids[20])+'/edit',{'status':'Walkout','walkout_reason':'Budget','competitor_shop':'','return_to':target})
            self.assertIn('walkout_page=2',response[2])
        finally:
            with closing(sqlite3.connect(str(self.dbpath))) as db,db:
                placeholders=','.join('?' for _ in ids)
                db.execute('DELETE FROM cre_walkout_records WHERE visit_id IN ('+placeholders+')',ids)
                db.execute('DELETE FROM customer_visits WHERE id IN ('+placeholders+')',ids)
    def test_configuration_and_field_restore(self):
        self.login()
        self.request('/cre/products',{'name':'Test Pendant'})
        self.assertTrue(self.sql("SELECT id FROM customer_product_types WHERE name='Test Pendant'"))
        self.request('/cre/sales-people',{'name':'New Sales Test','store_id':1})
        self.request('/cre/managers',{'name':'New Manager Test','store_id':1})
        self.assertEqual(self.sql("SELECT count(*) FROM cre_staff_members WHERE name LIKE 'New % Test'")[0][0],2)
        self.request('/cre/fields/1/remove',{})
        self.assertEqual(self.sql('SELECT is_active FROM customer_form_fields WHERE id=1')[0][0],0)
        self.request('/cre/fields/1/restore',{})
        self.assertEqual(self.sql('SELECT is_active FROM customer_form_fields WHERE id=1')[0][0],1)

    def test_super_admin_manages_stores_and_store_logins(self):
        self.login(user='executive')
        self.assertEqual(self.request('/cre/access')[0], 403)
        self.assertEqual(self.request('/cre/access/stores', {
            'store_name':'Blocked Store','city':'Indore','visit_prefix':'BLK',
            'store_address':'Address','timings':'10-8'
        })[0], 403)
        self.request('/cre/logout')
        self.login()
        page = self.request('/cre/access')
        self.assertEqual(page[0], 200)
        self.assertIn(b'Stores &amp; Logins', page[1])
        self.request('/cre/access/stores', {
            'store_name':'Second Test Store','city':'Bhopal','visit_prefix':'BTS',
            'store_address':'Test address 2','timings':'11-7'
        })
        store_id = self.sql("SELECT id FROM stores WHERE store_name='Second Test Store'")[0][0]
        self.request('/cre/access/accounts', {
            'username':'second_login','display_name':'Second CRE','password':'Second-password-123','store_id':store_id
        })
        account = self.sql("SELECT id,role,store_id FROM admin_users WHERE username='second_login'")[0]
        self.assertEqual(account[1:], ('customer_relation_executive', store_id))
        self.request('/cre/access/accounts/'+str(account[0])+'/update', {
            'display_name':'Updated CRE','password':'Updated-password-123','store_id':1
        })
        self.assertEqual(self.sql(
            "SELECT display_name,store_id FROM admin_users WHERE username='second_login'"
        ), [('Updated CRE', 1)])
        self.request('/cre/logout')
        self.assertEqual(self.login(user='second_login', password='Updated-password-123')[0], 200)

    def test_store_users_are_isolated_across_reads_writes_and_exports(self):
        self.login()
        first_name = 'Store isolation Indore'
        second_name = 'Store isolation Ujjain'
        forged_name = 'Forged store assignment'
        self.request('/cre/customers', self.form(
            customer_name=first_name, status='Walkout', walkout_reason='Budget', remark='Store isolation test'
        ))
        self.request('/cre/customers', self.form(
            customer_name=second_name, status='Walkout', walkout_reason='Budget', remark='Store isolation test', store_id=2,
            handled_by_id=3, assigned_manager_id=4,
        ))
        second_visit = self.sql(
            'SELECT id FROM customer_visits WHERE customer_name=?', (second_name,)
        )[0][0]
        admin_all = self.request('/cre/dashboard')[1]
        self.assertIn(first_name.encode(), admin_all)
        self.assertIn(second_name.encode(), admin_all)
        admin_filtered = self.request('/cre/dashboard?store_id=2')[1]
        self.assertNotIn(first_name.encode(), admin_filtered)
        self.assertIn(second_name.encode(), admin_filtered)
        city_report = self.request(
            '/cre/reports?date_from=2000-01-01&date_to=2100-01-01&city=Ujjain'
        )[1]
        self.assertIn(b'<option value="Ujjain" selected>Ujjain</option>', city_report)
        self.assertNotIn(first_name.encode(), city_report)
        self.assertIn(second_name.encode(), city_report)
        city_export = load_workbook(io.BytesIO(self.request(
            '/cre/reports/export/excel?date_from=2000-01-01&date_to=2100-01-01&city=Ujjain'
        )[1]))
        city_exported = ' '.join(
            str(cell.value or '') for sheet in city_export.worksheets for row in sheet for cell in row
        )
        self.assertNotIn(first_name, city_exported)
        self.assertIn(second_name, city_exported)

        self.request('/cre/logout')
        self.login(user='executive')
        dashboard = self.request('/cre/dashboard?store_id=2')[1]
        self.assertIn(first_name.encode(), dashboard)
        self.assertNotIn(second_name.encode(), dashboard)
        report = self.request('/cre/reports?date_from=2000-01-01&date_to=2100-01-01&store_id=2')[1]
        self.assertIn(first_name.encode(), report)
        self.assertNotIn(second_name.encode(), report)

        export = load_workbook(io.BytesIO(self.request(
            '/cre/reports/export/excel?date_from=2000-01-01&date_to=2100-01-01&store_id=2'
        )[1]))
        exported = ' '.join(
            str(cell.value or '') for sheet in export.worksheets for row in sheet for cell in row
        )
        self.assertIn(first_name, exported)
        self.assertNotIn(second_name, exported)

        response = self.request('/cre/customers', self.form(
            customer_name=forged_name, status='Draft', store_id=2,
            handled_by_id=1, assigned_manager_id=2,
        ))
        self.assertNotIn('error=', response[2])
        self.assertEqual(self.sql(
            'SELECT store_id FROM customer_visits WHERE customer_name=?', (forged_name,)
        ), [(1,)])

        self.assertIn('error=', self.request(
            '/cre/customers/'+str(second_visit)+'/status', {'status':'Billed'}
        )[2])
        self.assertEqual(self.sql(
            'SELECT status FROM customer_visits WHERE id=?', (second_visit,)
        ), [('Walkout',)])
        followup = self.request(
            '/cre/reports/walkouts/'+str(second_visit)+'/followups/1',
            {'note':'Cross-store note','return_to':'/cre/walkout-register'},
        )
        self.assertIn('error=', followup[2])
        self.assertEqual(self.sql(
            'SELECT count(*) FROM cre_walkout_followups WHERE visit_id=?', (second_visit,)
        ), [(0,)])
        self.assertEqual(self.request('/cre/products')[0], 200)
        self.assertIn(b'Sales People', dashboard)
        self.assertIn(b'Managers', dashboard)
        self.assertIn(b'Products', dashboard)
        self.assertIn(b'Walkout &amp; Competitors', dashboard)
        self.request('/cre/sales-people', {'name':'Store Scoped Sales','store_id':2})
        self.assertEqual(self.sql(
            "SELECT store_id FROM cre_staff_members WHERE name='Store Scoped Sales'"
        ), [(1,)])
        self.request('/cre/products', {'name':'Store User Product'})
        self.assertEqual(self.sql(
            "SELECT count(*) FROM customer_product_types WHERE name='Store User Product'"
        ), [(1,)])

        self.request('/cre/logout')
        self.login(user='second_executive')
        second_dashboard = self.request('/cre/dashboard')[1]
        self.assertIn(second_name.encode(), second_dashboard)
        self.assertNotIn(first_name.encode(), second_dashboard)

if __name__=='__main__': unittest.main()
